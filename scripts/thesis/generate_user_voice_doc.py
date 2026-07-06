"""Generate the User-Voice Requirements Word document for the thesis.

Output: thesis_user_voice_requirements.docx in the project root.

Layout matches the FR / NFR / Constraints docs: one continuous table with
five columns (ID, User Type, Action, Result, Success Measure). Role names
appear as merged-cell section headers between user-story groups.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# Each entry: (action, result, success_measure)
ROLES = [
    (
        'Compliance Analyst',
        [
            (
                'upload an internal BBK policy and receive a structured coverage analysis against the regulatory corpus of Bahrain, India, and Kuwait',
                'I can identify compliance gaps without manually reading every regulation',
                'analysis completion within 90 seconds for a typical policy document',
            ),
            (
                'see every finding in a generated report include a citation back to the exact regulatory clause it relies on',
                'I can verify the system’s reasoning and present audit-defensible evidence',
                '100% of findings carrying a clause-level citation traceable to the source document',
            ),
            (
                'filter retrieval results by jurisdiction, regulator, and document type',
                'I can focus my analysis on the regulatory regime relevant to a specific compliance question',
                'retrieval response time of under 2 seconds on a warm cache',
            ),
            (
                'export analysis reports in XLSX, DOCX, and PDF formats',
                'I can share findings with internal stakeholders using their preferred working format',
                'all three export formats being downloadable from a single report view',
            ),
        ],
    ),
    (
        'Legal Reviewer',
        [
            (
                'see each AI-generated finding flagged with a confidence score and a list of supporting evidence chunks',
                'I can prioritise which findings require deeper human review',
                'every finding carrying both a confidence value and at least one linked evidence chunk',
            ),
            (
                'approve, reject, or annotate AI-generated findings',
                'the system’s outputs become an auditable record of human-validated compliance positions',
                'reviewer actions being persisted to the audit log within the same request',
            ),
            (
                'have findings where the AI flagged a potential hallucination or prompt-injection signal routed into a quarantine queue',
                'suspect content does not reach decision-makers without explicit review',
                '100% of flagged items being held until a reviewer dispositions them',
            ),
            (
                'compare a policy clause side-by-side against the equivalent clauses across all three jurisdictions',
                'I can defensibly decide whether coverage is equivalent, stronger, weaker, partial, or absent',
                'the comparison view rendering all available cross-jurisdiction analogues in a single screen',
            ),
        ],
    ),
    (
        'Administrator',
        [
            (
                'view a system-health dashboard summarising LLM provider status, vector store size, lexical index health, active sessions, and recent reasoning failures',
                'I can spot operational issues before they impact analysts',
                'the dashboard refreshing on demand and reflecting the last 24 hours of activity',
            ),
            (
                'assign and revoke user roles (Analyst, Reviewer, Administrator)',
                'access to sensitive analysis capabilities is bounded by the user’s responsibility',
                'every role change producing an audit-log entry identifying the actor, target, before-state, and after-state',
            ),
            (
                'swap the active LLM provider via configuration without redeploying the application',
                'the system can move between managed cloud and on-premise inference as procurement, cost, or data-sovereignty constraints change',
                'the change taking effect after a configuration reload with no code modification',
            ),
            (
                'monitor bias parity across user roles and jurisdictions',
                'I can detect when approval rates or confidence scores diverge in ways that suggest systemic bias',
                'configurable threshold breaches surfacing as anomaly alerts on the monitoring dashboard',
            ),
        ],
    ),
]


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def _set_run_style(run, *, bold=False, color=NAVY, size=10):
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.size = Pt(size)


def _style_paragraph(paragraph, *, bold=False, color=NAVY, size=10, align=None):
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        _set_run_style(run, bold=bold, color=color, size=size)


def main():
    base = Path(__file__).resolve().parent.parent
    out_path = base / 'thesis_user_voice_requirements.docx'
    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = base / f'thesis_user_voice_requirements_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    doc = Document()

    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        new_width, new_height = section.page_height, section.page_width
        section.page_width = new_width
        section.page_height = new_height
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    title = doc.add_heading('User-Voice Requirements', level=0)
    for run in title.runs:
        run.font.color.rgb = NAVY

    intro = doc.add_paragraph(
        "This section restates the system’s primary capabilities in user-voice "
        "format, grouped by the three operational roles: Compliance Analyst, Legal "
        "Reviewer, and Administrator. Each user story follows the structure “As "
        "a [user type], I want to [action], so that [result],” paired with a "
        "concrete success measure that defines when the requirement is satisfied. "
        "These user stories complement the atomic functional requirements in "
        "Section 3.1.2 by expressing them from the perspective of the end users "
        "who will operate the system."
    )
    for run in intro.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY

    total_rows = 1 + sum(1 + len(stories) for _, stories in ROLES)
    table = doc.add_table(rows=total_rows, cols=5)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    widths = [
        Cm(1.3),   # ID
        Cm(3.2),   # User Type
        Cm(7.0),   # Action
        Cm(7.0),   # Result
        Cm(5.5),   # Success Measure
    ]
    for col_idx, width in enumerate(widths):
        for cell in table.columns[col_idx].cells:
            cell.width = width

    # Header row.
    header = table.rows[0]
    headers = ['ID', 'As a… (User Type)', 'I want to… (Action)',
               'So that… (Result)', 'Measured by… (Success Measure)']
    for i, h in enumerate(headers):
        header.cells[i].text = h
        shade_cell(header.cells[i], '002583')
        header.cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = header.cells[i].paragraphs[0]
        para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if i == 0
                          else WD_ALIGN_PARAGRAPH.LEFT)
        for run in para.runs:
            _set_run_style(run, bold=True, color=WHITE, size=10)

    usr_counter = 0
    row_idx = 1
    for role_name, stories in ROLES:
        # Role divider row: merge all five cells, navy fill, white bold text.
        role_row = table.rows[row_idx]
        merged = role_row.cells[0]
        for c in range(1, 5):
            merged = merged.merge(role_row.cells[c])
        merged.text = role_name
        shade_cell(merged, '002583')
        merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = merged.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            _set_run_style(run, bold=True, color=WHITE, size=11)
        row_idx += 1

        for action, result, measure in stories:
            usr_counter += 1
            data_row = table.rows[row_idx]
            data_row.cells[0].text = f'USR-{usr_counter:02d}'
            data_row.cells[1].text = role_name
            data_row.cells[2].text = action
            data_row.cells[3].text = result
            data_row.cells[4].text = measure
            for i, cell in enumerate(data_row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                for paragraph in cell.paragraphs:
                    align = (WD_ALIGN_PARAGRAPH.CENTER if i == 0
                             else WD_ALIGN_PARAGRAPH.LEFT)
                    _style_paragraph(paragraph, color=NAVY, size=10, align=align)
            # Bold the ID and the User Type.
            for run in data_row.cells[0].paragraphs[0].runs:
                run.font.bold = True
            for run in data_row.cells[1].paragraphs[0].runs:
                run.font.bold = True
            row_idx += 1

    # ── Table caption (below table, matching thesis convention) ──
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run('Table 3. User-Voice Requirements by Role')
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = NAVY
    run.font.size = Pt(11)

    doc.save(str(out_path))
    print(f'Wrote: {out_path} ({usr_counter} user stories across {len(ROLES)} roles)')


if __name__ == '__main__':
    main()
