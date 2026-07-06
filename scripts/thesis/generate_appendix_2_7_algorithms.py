"""Generate Appendix 2.7 — 12 algorithm pseudocode boxes.

Output: thesis_docs/appendix_2_7_algorithms.docx

A2.7.1 Auxiliary algorithms (9):
  StrictnessScore, DivergenceRanking, RiskScore, QueryRouter,
  TermDictionaryExpand, AutoRouteScope, ChunkClassify, SafeFallback,
  InvokeLocalLLM

A2.7.2 Django integration algorithms (3):
  SyncGapOnApproval, LaunchSubprocess, CascadeApproval

Each algorithm is rendered as a bordered box with monospace numbered
steps, Input/Output declarations, and a centred caption underneath.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLACK = RGBColor(0x00, 0x00, 0x00)
NAVY = BLACK
MUTED = BLACK
HEADER_FILL = 'D9D9D9'
BOX_BORDER = '1F2937'


def _set_cell_borders(cell, top=True, bottom=True, left=True, right=True,
                       color=BOX_BORDER, sz=6):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcBorders = tc_pr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tc_pr.append(tcBorders)
    for side, enabled in (('top', top), ('bottom', bottom),
                          ('left', left), ('right', right)):
        existing = tcBorders.find(qn(f'w:{side}'))
        if existing is not None:
            tcBorders.remove(existing)
        border = OxmlElement(f'w:{side}')
        if enabled:
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), str(sz))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
        else:
            border.set(qn('w:val'), 'nil')
        tcBorders.append(border)


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_run_style(run, *, bold=False, italic=False, color=BLACK, size=10,
                  font='Consolas'):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)
    run.font.name = font


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 2', 2: 'Heading 3', 3: 'Heading 4'}
    size_map = {1: 14, 2: 12, 3: 11}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_paragraph(doc, text, *, size=11, italic=False, color=BLACK,
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    return p


def add_caption(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    r1 = p.add_run(label + ' ')
    r1.font.bold = True
    r1.font.color.rgb = NAVY
    r1.font.size = Pt(10)
    r1.font.name = 'Calibri'
    r2 = p.add_run(text)
    r2.font.italic = True
    r2.font.color.rgb = MUTED
    r2.font.size = Pt(10)
    r2.font.name = 'Calibri'


def add_algorithm_box(doc, name, inputs, outputs, steps):
    """Render one algorithm as a bordered single-cell table with monospace
    steps inside."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell in table.columns[0].cells:
        cell.width = Cm(16.0)

    cell = table.rows[0].cells[0]
    _set_cell_borders(cell, color=BOX_BORDER, sz=6)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    # Clear default empty paragraph, we'll write our own
    cell.text = ''

    # Algorithm name
    p_name = cell.paragraphs[0]
    p_name.paragraph_format.space_after = Pt(4)
    r = p_name.add_run(f'Algorithm — {name}')
    set_run_style(r, bold=True, size=11, font='Calibri')

    # Input line
    p_in = cell.add_paragraph()
    p_in.paragraph_format.space_after = Pt(2)
    r1 = p_in.add_run('Input:   ')
    set_run_style(r1, bold=True, size=10)
    r2 = p_in.add_run(inputs)
    set_run_style(r2, size=10)

    # Output line
    p_out = cell.add_paragraph()
    p_out.paragraph_format.space_after = Pt(8)
    r3 = p_out.add_run('Output:  ')
    set_run_style(r3, bold=True, size=10)
    r4 = p_out.add_run(outputs)
    set_run_style(r4, size=10)

    # Steps
    for i, step in enumerate(steps, start=1):
        ps = cell.add_paragraph()
        ps.paragraph_format.left_indent = Cm(0.4)
        ps.paragraph_format.space_after = Pt(1)
        r = ps.add_run(f'{i:>2}. {step}')
        set_run_style(r, size=10)


# -------------------------------------------------------------------------
# ALGORITHM DATA
# -------------------------------------------------------------------------

# A2.7.1 — Auxiliary algorithms

ALG_STRICTNESS_SCORE = {
    'name': 'StrictnessScore',
    'inputs': 'regulation chunk c, jurisdiction j',
    'outputs': 'strictness score s in [0, 1]',
    'steps': [
        'tokens   <- normalise(c.text)',
        'penalty  <- match_weight(tokens, PENALTY_LEXICON)',
        'enforce  <- match_weight(tokens, ENFORCEMENT_LEXICON)',
        'coverage <- match_weight(tokens, COVERAGE_LEXICON)',
        'w        <- JURISDICTION_WEIGHTS[j]',
        's        <- w.penalty * penalty',
        '            + w.enforce * enforce',
        '            + w.coverage * coverage',
        'return clamp(s, 0, 1)',
    ],
}

ALG_DIVERGENCE_RANKING = {
    'name': 'DivergenceRanking',
    'inputs': 'list of clause pairs P = [(a_i, b_i)]',
    'outputs': 'P ordered by divergence, highest first',
    'steps': [
        'for each (a, b) in P do',
        '    s_a <- StrictnessScore(a, a.jurisdiction)',
        '    s_b <- StrictnessScore(b, b.jurisdiction)',
        '    sim <- cosine(embed(a), embed(b))',
        '    div <- w1 * abs(s_a - s_b)',
        '            + w2 * (1 - sim)',
        '            + w3 * verdict_weight(a, b)',
        '    pair.divergence <- div',
        'return sort(P, key = divergence, desc)',
    ],
}

ALG_RISK_SCORE = {
    'name': 'RiskScore',
    'inputs': 'gap g with severity, jurisdiction, topic',
    'outputs': 'risk score r in [0, 1]',
    'steps': [
        'sev   <- SEVERITY_WEIGHT[g.severity]',
        'jur   <- JURISDICTION_RISK[g.jurisdiction]',
        'topic <- TOPIC_IMPACT[g.topic]',
        'cov   <- COVERAGE_FACTOR[g.coverage_status]',
        'r     <- 0.40 * sev',
        '          + 0.25 * jur',
        '          + 0.25 * topic',
        '          + 0.10 * cov',
        'return clamp(r, 0, 1)',
    ],
}

ALG_QUERY_ROUTER = {
    'name': 'QueryRouter',
    'inputs': 'user query q',
    'outputs': 'route in {regulatory, approved, general}',
    'steps': [
        'tokens <- normalise(q)',
        'if matches_any(tokens, APPROVED_KEYWORDS) then',
        '    return "approved"',
        'if matches_any(tokens, REGULATORY_KEYWORDS) then',
        '    return "regulatory"',
        'score <- classifier_confidence(q)',
        'if score.regulatory > 0.6 then',
        '    return "regulatory"',
        'return "general"',
    ],
}

ALG_TERM_DICTIONARY_EXPAND = {
    'name': 'TermDictionaryExpand',
    'inputs': 'query q, term dictionary D (~200 synonyms)',
    'outputs': 'expanded query q_exp',
    'steps': [
        'q_exp <- q',
        'for each term t in tokens(q) do',
        '    if t in D.keys then',
        '        for each syn in D[t] do',
        '            q_exp <- q_exp + " OR " + syn',
        'return q_exp',
    ],
}

ALG_AUTO_ROUTE_SCOPE = {
    'name': 'AutoRouteScope',
    'inputs': 'query q, available regulations R',
    'outputs': 'filtered regulations R_scope',
    'steps': [
        'topics <- detect_topics(q, TOPIC_TAXONOMY)',
        'if topics is empty then',
        '    return R                  // no scoping',
        'R_scope <- []',
        'for each r in R do',
        '    if r.topics intersects topics then',
        '        R_scope.append(r)',
        'return R_scope',
    ],
}

ALG_CHUNK_CLASSIFY = {
    'name': 'ChunkClassify',
    'inputs': 'chunk c',
    'outputs': 'tag in chunk_tags taxonomy or "unknown"',
    'steps': [
        'rule_tag <- rule_match(c.text, RULE_TAG_PATTERNS)',
        'if rule_tag is not null then',
        '    return rule_tag',
        'prompt   <- build_classify_prompt(c, TAXONOMY)',
        'response <- InvokeLocalLLM(prompt)',
        'tag      <- parse_tag(response)',
        'if tag not in TAXONOMY or confidence < 0.6 then',
        '    return "unknown"',
        'return tag',
    ],
}

ALG_SAFE_FALLBACK = {
    'name': 'SafeFallback',
    'inputs': 'schema T (Pydantic model), reason r',
    'outputs': 'typed-empty instance of T',
    'steps': [
        'instance <- T()                 // all defaults',
        'instance.confidence <- 0.3',
        'instance.citation_verified <- False',
        'instance.hallucination_risk <- 1.0',
        'instance.notes <- "SafeFallback: " + r',
        'log_event(reasoning.validation_error, reason = r)',
        'return instance',
    ],
}

ALG_INVOKE_LOCAL_LLM = {
    'name': 'InvokeLocalLLM',
    'inputs': 'prompt p, timeout t (default 30 s)',
    'outputs': 'LLM response or raises LLMError',
    'steps': [
        'attempt <- 0',
        'while attempt < MAX_RETRIES do',
        '    try',
        '        return ollama_client.generate(p, timeout = t)',
        '    catch TimeoutError, ConnectionError',
        '        attempt <- attempt + 1',
        '        sleep(backoff(attempt))',
        'raise LLMError("Local LLM unavailable after retries")',
    ],
}


# A2.7.2 — Django integration algorithms

ALG_SYNC_GAP_ON_APPROVAL = {
    'name': 'SyncGapOnApproval',
    'inputs': 'ObligationMapping o (post-save signal payload)',
    'outputs': 'Gap row inserted or updated',
    'steps': [
        'if o.lifecycle is not APPROVED then',
        '    return',
        'if o.coverage in {covered, full} then',
        '    Gap.objects.filter(obligation = o).delete()',
        '    return',
        'gap, created <- Gap.objects.update_or_create(',
        '    obligation = o,',
        '    defaults = {',
        '        "severity": o.derived_severity,',
        '        "risk":     RiskScore(o),',
        '        "due_date": SEVERITY_TO_DUE_DATE[o.severity]',
        '    })',
        'log_event(mapping.gap_synced, target = gap)',
    ],
}

ALG_LAUNCH_SUBPROCESS = {
    'name': 'LaunchSubprocess',
    'inputs': 'job kind k, parent pk',
    'outputs': 'subprocess PID',
    'steps': [
        'parent <- Model.objects.get(pk = parent_pk)',
        'parent.status <- QUEUED',
        'parent.save()',
        'cmd <- [PYTHON, MANAGE_PY, "run_" + k + "_job",',
        '        "--pk", str(parent_pk)]',
        'env <- os.environ.copy()',
        'env["DJANGO_SUBPROCESS"] <- "1"',
        'p   <- subprocess.Popen(cmd, env = env)',
        'log_event(k + ".run", target = parent)',
        'return p.pid',
    ],
}

ALG_CASCADE_APPROVAL = {
    'name': 'CascadeApproval',
    'inputs': 'parent analysis A',
    'outputs': 'approved children count',
    'steps': [
        'if A.lifecycle is not APPROVED then',
        '    return 0',
        'n <- 0',
        'for each child c in A.children do',
        '    if c.lifecycle is REJECTED then',
        '        continue              // explicit reject wins',
        '    if c.lifecycle in {DRAFT, REVIEWED} then',
        '        c.lifecycle <- APPROVED',
        '        c.human_override <- True',
        '        c.save()',
        '        log_event(review.cascade_approve, target = c)',
        '        n <- n + 1',
        'return n',
    ],
}


# Display labels (A.7.1.a, etc.)
AUXILIARY = [
    ('A.7.1.a', ALG_STRICTNESS_SCORE),
    ('A.7.1.b', ALG_DIVERGENCE_RANKING),
    ('A.7.1.c', ALG_RISK_SCORE),
    ('A.7.1.d', ALG_QUERY_ROUTER),
    ('A.7.1.e', ALG_TERM_DICTIONARY_EXPAND),
    ('A.7.1.f', ALG_AUTO_ROUTE_SCOPE),
    ('A.7.1.g', ALG_CHUNK_CLASSIFY),
    ('A.7.1.h', ALG_SAFE_FALLBACK),
    ('A.7.1.i', ALG_INVOKE_LOCAL_LLM),
]

DJANGO_INTEGRATION = [
    ('A.7.2.a', ALG_SYNC_GAP_ON_APPROVAL),
    ('A.7.2.b', ALG_LAUNCH_SUBPROCESS),
    ('A.7.2.c', ALG_CASCADE_APPROVAL),
]


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / 'appendix_2_7_algorithms.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'appendix_2_7_algorithms_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_heading(doc, 'Appendix 2.7 — Algorithm reference', level=1)
    add_paragraph(
        doc,
        'This appendix expands the auxiliary and Django-integration '
        'algorithms named but not boxed in §3.2.3. A2.7.1 lists nine '
        'auxiliary procedures called by the main reasoning, retrieval, '
        'and ingestion paths. A2.7.2 lists three patterns that connect '
        'the AI pipeline to the Django web layer. Weights, thresholds, '
        'and lookup tables referenced in the pseudocode are defined in '
        'Appendix 2.8.',
        size=11,
    )

    add_heading(doc, 'A2.7.1 Auxiliary algorithms', level=2)
    add_paragraph(
        doc,
        'Nine procedures expanded here as numbered pseudocode boxes. '
        'Each box names the inputs and outputs followed by the step '
        'sequence in execution order.',
        size=11,
    )
    for label, alg in AUXILIARY:
        add_algorithm_box(doc,
                          alg['name'],
                          alg['inputs'],
                          alg['outputs'],
                          alg['steps'])
        add_caption(doc, f'Algorithm {label}.', alg['name'])

    add_heading(doc, 'A2.7.2 Django integration algorithms', level=2)
    add_paragraph(
        doc,
        'Three patterns that glue the AI pipeline to the Django web '
        'layer.',
        size=11,
    )
    for label, alg in DJANGO_INTEGRATION:
        add_algorithm_box(doc,
                          alg['name'],
                          alg['inputs'],
                          alg['outputs'],
                          alg['steps'])
        add_caption(doc, f'Algorithm {label}.', alg['name'])

    doc.save(out_path)
    total = len(AUXILIARY) + len(DJANGO_INTEGRATION)
    print(f'Wrote {total} algorithm boxes to {out_path}')


if __name__ == '__main__':
    main()
