"""Generate Appendix 3 Implementation Supplements as a Word document.

Output: thesis_docs/appendix_3_implementation.docx

Eleven reference cards. Each card covers one implementation phase that
was either not deep enough to justify main-body real estate, or whose
detail belongs out of the chapter flow. Format per card:
  - Purpose          (1 sentence)
  - Tools / libraries (bulleted line)
  - Setup commands    (monospace block; copy-pasteable)
  - Code outline      (1 PNG where applicable)
  - Verification      (1 line)

Figure numbering: A.3.1 onward. The main body owns figures 24 through 54.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)


# ─────────────────────────── Style helpers ────────────────────────────────

def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_run_style(run, *, bold=False, italic=False, color=NAVY, size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 2', 2: 'Heading 3', 3: 'Heading 4'}
    size_map = {1: 14, 2: 12, 3: 11}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_caption(doc, label: str, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    r1 = p.add_run(label + ' ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run(caption)
    set_run_style(r2, italic=True, color=MUTED, size=10)


def add_image(doc, image_path: Path, width_cm: float = 15.5):
    if not image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'[MISSING IMAGE: {image_path.name}]')
        set_run_style(run, bold=True, italic=True, color=MUTED, size=10)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def add_monospace(doc, text: str, *, size: int = 9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(size)
    run.font.color.rgb = DARK_GRAY


def add_label_line(doc, label: str, value: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(label + ' ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run(value)
    set_run_style(r2, color=DARK_GRAY, size=10)


# ───────────────────────────── Document body ──────────────────────────────

def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / 'appendix_3_implementation.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'appendix_3_implementation_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    code = base / 'diagrams' / 'code_3_3'
    shots = base / 'diagrams' / 'screenshots_3_3'

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_heading(doc, 'Appendix 3 Implementation Supplements', level=1)
    add_paragraph(
        doc,
        'This appendix supports §3.3 with the operational detail every '
        'phase would otherwise carry in the main body. Each card lists the '
        'phase purpose, the tools and libraries, the commands a '
        'replicator runs, the code outline, and the verification step '
        'that confirms the phase landed.',
    )

    # ─── A.3.1 Host OS preparation ────────────────────────────────────────
    add_heading(doc, 'A.3.1 Host operating system preparation', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Prepare a clean Windows 11 host with Python, Git, and Visual '
        'Studio Code so the rest of the build runs without environment '
        'surprises.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'Windows 11 Pro; Python 3.10 or newer; Git for Windows; '
        'Visual Studio Code with Python, Pylance, SQLite Viewer, '
        'and CodeSnap extensions.',
    )
    add_label_line(doc, 'Setup commands:', '')
    add_monospace(doc, (
        '> winget install Python.Python.3.12\n'
        '> winget install Git.Git\n'
        '> winget install Microsoft.VisualStudioCode\n'
        '> code --install-extension ms-python.python\n'
        '> code --install-extension ms-python.vscode-pylance\n'
        '> code --install-extension alexcvzz.vscode-sqlite\n'
        '> code --install-extension adpyke.codesnap\n'
        '> python --version\n'
        '> git --version\n'
    ))
    add_label_line(
        doc, 'Verification:',
        'The "python --version" command prints 3.10 or higher and Visual '
        'Studio Code lists the four extensions under installed.',
    )

    # ─── A.3.2 Project bootstrap ──────────────────────────────────────────
    add_heading(doc, 'A.3.2 Project bootstrap', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Clone the repository, create the virtual environment, and '
        'install pinned dependencies plus the platform-correct PyTorch '
        'build.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'Python venv; pip; the dependencies listed in requirements.txt '
        '(chromadb, llama-index-core, sentence-transformers, '
        'rank-bm25, langchain-text-splitters, docling, pymupdf, '
        'python-docx, beautifulsoup4, pandas, openpyxl, reportlab, '
        'requests); PyTorch installed from its own distribution channel.',
    )
    add_label_line(doc, 'Setup commands:', '')
    add_monospace(doc, (
        '> git clone <repository-url> cjpca-project\n'
        '> cd cjpca-project\n'
        '> python -m venv .venv\n'
        '> .venv\\Scripts\\activate\n'
        '> python -m pip install --upgrade pip\n'
        '> python -m pip install -r requirements.txt\n'
        '> python -m pip install torch --index-url https://download.pytorch.org/whl/cpu\n'
        '> copy .env.example .env\n'
    ))
    add_image(doc, code / 'AP01_requirements.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.1.', 'Requirements highlights')
    add_label_line(
        doc, 'Verification:',
        '"python -c \\"import chromadb, langchain, torch\\"" exits cleanly '
        'and the .env file lists the local-LLM endpoint and security '
        'toggles.',
    )

    # ─── A.3.3 Django scaffold and database init ──────────────────────────
    add_heading(doc, 'A.3.3 Django scaffold and database initialisation',
                level=2)
    add_label_line(
        doc, 'Purpose:',
        'Apply migrations for the contributed Django schemas, the '
        'two-factor auth schemas, and the ten project apps, and create '
        'the initial admin account.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'Django 5; SQLite with WAL journaling; django-otp; '
        'django-two-factor-auth.',
    )
    add_label_line(doc, 'Setup commands:', '')
    add_monospace(doc, (
        '> cd cjpca\n'
        '> python manage.py makemigrations\n'
        '> python manage.py migrate\n'
        '> python manage.py createsuperuser\n'
        '> python -c "from django.db import connection; '
        'connection.cursor().execute(\\\"PRAGMA journal_mode=WAL\\\")"\n'
    ))
    add_label_line(
        doc, 'Verification:',
        '"python manage.py migrate" prints "No migrations to apply" on '
        'second run; the admin user can sign in at /admin/ after MFA '
        'enrolment.',
    )

    # ─── A.3.4 Data gathering and corpus assembly ─────────────────────────
    add_heading(doc, 'A.3.4 Data gathering and corpus assembly', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Acquire the three jurisdictional regulations and the BBK internal '
        'policies, organise them on disk, and record provenance for each '
        'document.',
    )
    add_label_line(
        doc, 'Sources:',
        'Bahrain PDPL and implementing orders from the Legal Affairs '
        'Bureau portal; India Digital Personal Data Protection Act 2023 '
        'from the Ministry of Electronics and Information Technology '
        'gazette; Kuwait Data Privacy Protection Regulation from CITRA; '
        'BBK internal policies transferred under non-disclosure following '
        'a BBK consultation.',
    )
    add_label_line(doc, 'Storage layout:', '')
    add_monospace(doc, (
        'data/\n'
        '├── regulations/\n'
        '│   ├── bahrain/        PDPL + Orders (PDF)\n'
        '│   ├── india/          DPDP Act 2023 (PDF)\n'
        '│   └── kuwait/         DPPR (PDF)\n'
        '└── policies/           BBK internal policies (DOCX)\n'
    ))
    add_label_line(
        doc, 'Provenance fields per document:',
        'Source URL or document reference; retrieval date; SHA-256 '
        'content hash; jurisdiction tag; document type (regulation or '
        'policy); language verification (English only per NFR-21).',
    )
    add_label_line(doc, 'Pre-processing:', '')
    add_monospace(doc, (
        '> python manage.py sync_documents\n'
        '> python manage.py classify_doc_topics\n'
    ))
    add_image(doc, shots / 'AP_data_folder_full.png', width_cm=15.0)
    add_caption(doc, 'Figure A.3.2.', 'Corpus folder tree')
    add_label_line(
        doc, 'Verification:',
        'Each Document row in the Django admin lists a non-empty '
        'content_hash, jurisdiction, and classified_topics field.',
    )

    # ─── A.3.5 Model and index initialisation ─────────────────────────────
    add_heading(doc, 'A.3.5 Model and index initialisation', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Warm the HuggingFace cache with the three required models and '
        'create the ChromaDB collection plus the SQLite FTS5 virtual '
        'table.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'sentence-transformers (BGE-small embedder); cross-encoder/ms-'
        'marco-MiniLM (reranker); cross-encoder NLI model (verifier); '
        'chromadb persistent client; SQLite FTS5.',
    )
    add_label_line(doc, 'Setup commands:', '')
    add_monospace(doc, (
        '> python -c "from sentence_transformers import SentenceTransformer; '
        'SentenceTransformer(\\\"BAAI/bge-small-en-v1.5\\\")"\n'
        '> python manage.py shell -c "from retrieval.retriever import '
        '_get_service; _get_service()._get_index()"\n'
        '> python manage.py shell -c "from retrieval.bm25_store import _init; _init()"\n'
    ))
    add_image(doc, code / 'AP02_chromadb_init.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.3.', 'ChromaDB persistent client init')
    add_image(doc, code / 'AP03_fts5_ddl.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.4.', 'SQLite FTS5 virtual table DDL')
    add_image(doc, shots / 'AP_hf_cache_folder.png', width_cm=15.0)
    add_caption(doc, 'Figure A.3.5.', 'HuggingFace cache after warm')
    add_label_line(
        doc, 'Verification:',
        'The hf_cache directory contains the three model snapshots; '
        'chroma_data exists; "python -c \\"import sqlite3; '
        'sqlite3.connect(\\\"cjpca/db.sqlite3\\\").execute(\\\"SELECT name '
        'FROM sqlite_master WHERE name=\\\"bm25_index\\\"\\\")\\"" '
        'returns one row.',
    )

    # ─── A.3.6 Local LLM setup ────────────────────────────────────────────
    add_heading(doc, 'A.3.6 Local LLM setup', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Install the Ollama daemon, pull the default reasoning model, '
        'and confirm the provider abstraction is wired through the '
        'environment variables.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'Ollama (Windows installer from ollama.com); langchain-ollama; '
        'the env-var-switchable provider in reasoning/generator.py.',
    )
    add_label_line(doc, 'Setup commands:', '')
    add_monospace(doc, (
        '> winget install Ollama.Ollama\n'
        '> ollama serve              # start the daemon\n'
        '> ollama pull llama3.2:1b   # default model\n'
        '> ollama list               # confirm the model is present\n'
    ))
    add_image(doc, code / 'AP04_provider_abstraction.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.6.', 'Provider abstraction outline')
    add_image(doc, shots / 'AP_ollama_list.png', width_cm=15.0)
    add_caption(doc, 'Figure A.3.7.', 'Ollama list of pulled models')
    add_label_line(
        doc, 'Verification:',
        '"curl http://localhost:11434/api/tags" returns the pulled model; '
        'a comparison run from the application completes without '
        'falling through to SafeFallback.',
    )

    # ─── A.3.7 Workflow engines ───────────────────────────────────────────
    add_heading(doc, 'A.3.7 Workflow engines', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Wire the three domain workflows (comparison, mapping, gap) '
        'behind the web layer through subprocess dispatch and '
        'signal-based cascades.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'subprocess.Popen for fire-and-forget job dispatch; '
        'django.db.models.signals.post_save for the gap-sync cascade.',
    )
    add_image(doc, code / 'AP05_subprocess_dispatch.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.8.', 'Subprocess dispatch command outline')
    add_image(doc, code / 'AP06_sync_gap_signal.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.9.', 'Gap synchronisation post-save signal')
    add_label_line(
        doc, 'Verification:',
        'Triggering a comparison from the UI spawns one '
        '"run_comparison_job" process; approving a gap-flagged '
        'ObligationMapping inserts a Gap row through the signal.',
    )

    # ─── A.3.8 Web layer ──────────────────────────────────────────────────
    add_heading(doc, 'A.3.8 Web layer', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Compose the ten Django apps, mount HTMX progressive enhancement, '
        'and stream long-running ingestion progress through Django '
        'Channels.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'Django 5 (views, URL conf, templates); django-htmx; Alpine.js; '
        'Tailwind CSS (CDN); Chart.js (CDN); Django Channels with the '
        'in-memory layer.',
    )
    add_image(doc, code / 'AP07_htmx_polling.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.10.', 'HTMX progress polling endpoint outline')
    add_label_line(
        doc, 'Verification:',
        'The dashboard renders the coverage chart in under five hundred '
        'milliseconds; uploading a small document streams progress over '
        'the WebSocket within two seconds of each stage change.',
    )

    # ─── A.3.9 Cybersecurity build details ────────────────────────────────
    add_heading(doc, 'A.3.9 Cybersecurity build details', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Supplementary code outlines for the cybersecurity layer covered '
        'in §3.3.12. Each outline is an NDA-safe stub of the production '
        'class or function.',
    )
    add_image(doc, code / 'AP08_complexity_validator.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.11.', 'Custom ComplexityValidator')
    add_image(doc, code / 'AP09_idle_session.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.12.', 'IdleSessionTimeoutMiddleware')
    add_image(doc, code / 'AP10_force_mfa.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.13.', 'ForceMFAEnrollmentMiddleware')
    add_image(doc, code / 'AP11_geofence.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.14.', 'GeoFenceMiddleware')
    add_image(doc, code / 'AP12_audit_log_model.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.15.', 'AuditLog model definition')
    add_image(doc, code / 'AP13_log_event.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.16.', 'log_event helper')
    add_image(doc, code / 'AP14_role_required.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.17.', 'role_required decorator')
    add_label_line(
        doc, 'Verification:',
        '"python manage.py check --deploy" reports zero deployment '
        'issues; the audit log captures every authentication event; '
        'logging in as each role gates the expected views.',
    )

    # ─── A.3.10 Export and reporting ──────────────────────────────────────
    add_heading(doc, 'A.3.10 Export and reporting', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Render finished comparisons and mappings as PDF or XLSX reports '
        'with a SHA-256 chain prefix on every row.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'ReportLab (PDF); openpyxl (XLSX); hashlib (SHA-256).',
    )
    add_image(doc, code / 'AP15_export_builder.png', width_cm=15.5)
    add_caption(doc, 'Figure A.3.18.', 'PDF and XLSX export builders')
    add_label_line(
        doc, 'Verification:',
        'Exporting a finished comparison produces a file whose last '
        'column is a sixteen-character hexadecimal prefix that '
        're-computes identically from the live row.',
    )

    # ─── A.3.11 Deployment to BBK workstation ─────────────────────────────
    add_heading(doc, 'A.3.11 Deployment to BBK workstation', level=2)
    add_label_line(
        doc, 'Purpose:',
        'Promote the application from a development environment to the '
        'BBK on-premises workstation with HTTPS enforced, security '
        'headers tightened, and the optional GeoFence enabled.',
    )
    add_label_line(
        doc, 'Tools and libraries:',
        'A reverse proxy (Caddy or nginx) terminating TLS 1.3 on port '
        '443; MaxMind GeoLite2 country database; the four backup '
        'directories (db.sqlite3, chroma_data, media, logs).',
    )
    add_label_line(doc, 'Setup commands (reverse proxy excerpt):', '')
    add_monospace(doc, (
        'cjpca.bbk.local {\n'
        '    encode gzip\n'
        '    tls /etc/ssl/cjpca.crt /etc/ssl/cjpca.key\n'
        '    header Strict-Transport-Security "max-age=31536000;'
        ' includeSubDomains; preload"\n'
        '    reverse_proxy 127.0.0.1:8000\n'
        '}\n'
    ))
    add_label_line(doc, 'Production environment toggles:', '')
    add_monospace(doc, (
        'DEBUG=False\n'
        'SECRET_KEY=<rotate to a fresh 60-char value>\n'
        'REQUIRE_MFA=True\n'
        'GEOFENCE_ENABLED=True\n'
        'GEOFENCE_ALLOWED_COUNTRIES=BH,IN,KW,AE\n'
        'SECURE_SSL_REDIRECT=True\n'
    ))
    add_image(doc, shots / 'AP_reverse_proxy_config.png', width_cm=15.0)
    add_caption(doc, 'Figure A.3.19.', 'Reverse proxy configuration excerpt')
    add_image(doc, shots / 'AP_backup_folders.png', width_cm=15.0)
    add_caption(doc, 'Figure A.3.20.', 'Backup folders on the BBK host')
    add_label_line(
        doc, 'Verification:',
        '"python manage.py check --deploy" reports zero issues; an '
        'incognito browser session reaches the login page over HTTPS '
        'and is denied from an out-of-allowlist country.',
    )

    doc.save(out_path)
    print(f'Wrote: {out_path}')


if __name__ == '__main__':
    main()
