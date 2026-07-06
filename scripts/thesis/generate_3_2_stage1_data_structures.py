"""Generate Section 3.2.1 Data Structures (Stage 1) as a Word document.

Output: thesis_docs/3_2_1_data_structures.docx

Contents:
  1. Opening paragraph to Section 3.2 Solution Design
  2. Section 3.2.1 Data Structures heading
  3. Three-design-principles prose
  4. Table T1 — Data-Structure Selection Rationale (10 rows)
  5. Figure D1a — Entity-Relationship Diagram (embedded from diagrams/)
  6. Figure D1b — Non-Relational Data Structures (embedded from diagrams/)
  7. Deep-dive paragraph: LangGraph state + RRF
  8. Deep-dive paragraph: Audit-hash chain

Style follows the existing thesis-doc generators (navy headings, 2 cm margins,
Table Grid). Paste the output document straight into the thesis.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)


# Table T1 rows. Each: (number, structure, rationale).
# Code-location column removed: source code is not shared (NDA).
T1_ROWS = [
    ('1', 'BGE-small (384-dim) + ChromaDB HNSW',
     'Open-source, persistent, cosine similarity with metadata filters. '
     'Scales to 50,000 chunks.'),
    ('2', 'SQLite FTS5 BM25 inverted index',
     'Sub-millisecond response, and handles the legal acronyms that '
     'vector embeddings struggle with.'),
    ('3', 'Reciprocal Rank Fusion (RRF)',
     'Combines BM25 and vector ranks without any score calibration step.'),
    ('4', 'Two-level chunk schema',
     'Splits on article headers and then merges on semantic similarity, '
     'so chunks preserve citation granularity.'),
    ('5', 'Content-hash on chunk id',
     'SHA-256 detects duplicate chunks at ingest time without scanning '
     'the database.'),
    ('6', 'LangGraph state TypedDict',
     'Carries the explicit verify, correct, and retry control flow with '
     'bounded iteration.'),
    ('7', 'Pydantic v2 schemas',
     'Hard parse boundary at the LLM and Python interface, with an '
     'auto-fixing parser for malformed output.'),
    ('8', 'chunk_tags 12-topic taxonomy',
     'Cheap and deterministic intent classification that drives the '
     'auto-routing of policy mappings.'),
    ('9', 'ObligationMapping lifecycle FSM',
     'Audit trail through FSM transitions across the draft, reviewed, '
     'approved, and rejected states.'),
    ('10', 'Audit-hash chain (SHA-256)',
     'Detects silent post-signing edits to exported PDF and XLSX '
     'artefacts.'),
]


def shade_cell(cell, hex_color: str) -> None:
    """Set a cell's background fill."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_run_style(run, *, bold=False, italic=False, color=NAVY, size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_heading(doc, text, *, level=1):
    """Add a styled heading. level=1 -> H2, level=2 -> H3 (one below)."""
    heading_style = 'Heading 2' if level == 1 else 'Heading 3'
    h = doc.add_paragraph(style=heading_style)
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    return h


def add_figure_caption(doc, label: str, caption: str):
    """Add a centred figure caption like 'Figure D1a — Description'."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    r1 = p.add_run(label + ' ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run(caption)
    set_run_style(r2, italic=True, color=MUTED, size=10)


def add_image(doc, image_path: Path, width_cm: float = 16.0):
    """Insert an image, centred, with the given width."""
    if not image_path.exists():
        # If the file is missing, add a placeholder paragraph so the docx is
        # still well-formed and the reviewer sees what is missing.
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'[MISSING IMAGE: {image_path.name}]')
        set_run_style(run, bold=True, italic=True, color=MUTED, size=10)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def add_t1_table(doc):
    """Render Table T1 - Data-Structure Selection Rationale."""
    # caption above
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r1 = p.add_run('Table T1 ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run('— Data-Structure Selection Rationale')
    set_run_style(r2, italic=True, color=MUTED, size=10)

    table = doc.add_table(rows=1 + len(T1_ROWS), cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [Cm(0.8), Cm(5.4), Cm(9.8)]
    for col_idx, w in enumerate(widths):
        for cell in table.columns[col_idx].cells:
            cell.width = w

    headers = ['#', 'Structure', 'Rationale']
    header_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = h
        shade_cell(cell, '002583')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            set_run_style(run, bold=True, color=WHITE, size=10)

    for r_idx, (n, struct, why) in enumerate(T1_ROWS, start=1):
        row = table.rows[r_idx]
        row.cells[0].text = n
        row.cells[1].text = struct
        row.cells[2].text = why
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_style(
                        run,
                        bold=(c_idx == 0),
                        color=DARK_GRAY,
                        size=9,
                    )


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / '3_2_1_data_structures.docx'

    # Versioned filename if the existing file is locked open in Word.
    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'3_2_1_data_structures_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    diagrams_dir = base / 'diagrams'
    fig_d1a = diagrams_dir / 'Figure_D1a_ERD (1).drawio.png'
    fig_d1b = diagrams_dir / 'Figure_D1b_NonRelational_v2 (1).drawio.png'

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # ── Section opening (§3.2 lead paragraph) ────────────────────────────────
    add_heading(doc, '3.2 Solution Design', level=1)
    add_paragraph(
        doc,
        'The design applies Section 3.1’s requirements across five '
        'layers: ingestion, retrieval, reasoning, web, and copilot. '
        'Security and observability are cross-cutting. The full '
        'traceability matrix is in Appendix 2.1.',
    )

    # ── 3.2.1 Data Structures ────────────────────────────────────────────────
    add_heading(doc, '3.2.1 Data Structures', level=2)

    add_paragraph(
        doc,
        'Three principles shape the design. First, the five layers '
        'stay separate with clean interfaces, so each can change on '
        'its own. Second, the system grounds every answer in retrieved '
        'corpus passages rather than fine-tuning a model, because '
        'legal text needs exact-clause citations and the corpus '
        'changes too often for a fine-tune cycle to keep up. Third, '
        'every LLM response is treated as a draft that must be '
        'checked. If the check fails, the system returns a fixed '
        'backup answer.',
    )

    add_paragraph(
        doc,
        'Table T1 lists the data structures used on the system’s core '
        'path and why each was chosen.',
    )

    add_t1_table(doc)

    add_paragraph(
        doc,
        'The relational structures are the Django models that store '
        'documents, analyses, obligations, and the audit trail. Figure '
        'D1a shows them as an ER diagram. The remaining structures '
        'cannot be drawn this way, so Figure D1b shows them separately.',
    )

    # Figure D1a
    add_image(doc, fig_d1a, width_cm=16.0)
    add_figure_caption(
        doc,
        'Figure D1a.',
        'Entity-Relationship Diagram of the Relational Data Layer.',
    )

    # Figure D1b
    add_image(doc, fig_d1b, width_cm=14.0)
    add_figure_caption(
        doc,
        'Figure D1b.',
        'Non-Relational Data Structures.',
    )

    # Deep-dive 1: LangGraph state + RRF
    add_paragraph(
        doc,
        'Two structures deserve a closer look. The LangGraph reasoning '
        'state is a small typed dictionary that carries the request, '
        'retrieved chunks, reasoning trace, and a retry counter through '
        'every node of the verify, correct, and finalize graph. Because '
        'the counter is a plain number, the state machine can stop '
        'after a set number of attempts and return a fixed backup '
        'answer. Reciprocal Rank Fusion merges the top BM25 and '
        'ChromaDB results using score = sum(1 / (k + r)) with k = 60, '
        'with no calibration needed between the two scales.',
    )

    # Deep-dive 2: Audit-hash chain
    add_paragraph(
        doc,
        'The audit-hash chain makes tampering visible on exported '
        'files. Each PDF and XLSX export carries a sixteen-character '
        'SHA-256 prefix computed from the analysis’s identity and '
        'decision metadata. A recipient can recompute the hash and '
        'spot silent post-signing edits, without a full '
        'digital-signature setup. Appendix 2.3 carries the '
        'entity-relationship detail.',
    )

    doc.save(out_path)
    print(f'Wrote: {out_path}')


if __name__ == '__main__':
    main()
