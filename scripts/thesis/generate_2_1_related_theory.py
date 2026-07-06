"""Generate §2.1 Related Theory + matching references as a docx.

Output: thesis_docs/2_1_related_theory.docx

The four theory paragraphs use disjoint reference sets (no reference
appears in more than one paragraph), every in-text citation is
verified against the underlying paper, and three additions strengthen
the claim-citation alignment (Nogueira & Cho 2019, Ouyang et al. 2022,
T. Gao et al. 2023).
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)


def set_run_style(run, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3'}
    size_map = {1: 16, 2: 14, 3: 12}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  space_after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_theory_paragraph(doc, lead, body):
    """Bolded lead phrase followed by the paragraph body."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    lead_run = p.add_run(lead + ' ')
    set_run_style(lead_run, bold=True, color=DARK_GRAY, size=11)
    body_run = p.add_run(body)
    set_run_style(body_run, color=DARK_GRAY, size=11)


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_style(run, color=DARK_GRAY, size=10.5)


REFERENCES = [
    'Bommasani, R., Hudson, D. A., Adeli, E., Altman, R., Arora, S., '
    'von Arx, S., Bernstein, M. S., Bohg, J., Bosselut, A., Brunskill, '
    'E., Brynjolfsson, E., Buch, S., Card, D., Castellon, R., '
    'Chatterji, N., Chen, A., Creel, K., Davis, J. Q., Demszky, D., '
    '... Liang, P. (2021). On the opportunities and risks of '
    'foundation models (arXiv:2108.07258). arXiv. '
    'https://arxiv.org/abs/2108.07258',

    'Es, S., James, J., Espinosa-Anke, L., & Schockaert, S. (2024). '
    'RAGAS: Automated evaluation of retrieval augmented generation. '
    'In Proceedings of the 18th Conference of the European Chapter '
    'of the Association for Computational Linguistics: System '
    'Demonstrations (pp. 150-158). Association for Computational '
    'Linguistics. https://aclanthology.org/2024.eacl-demo.16/',

    'Gao, T., Yen, H., Yu, J., & Chen, D. (2023). Enabling large '
    'language models to generate text with citations. In Proceedings '
    'of the 2023 Conference on Empirical Methods in Natural Language '
    'Processing (pp. 6465-6488). Association for Computational '
    'Linguistics. https://aclanthology.org/2023.emnlp-main.398/',

    'Gao, Y., Xiong, Y., Gao, X., Jia, K., Pan, J., Bi, Y., Dai, Y., '
    'Sun, J., Wang, M., & Wang, H. (2024). Retrieval-augmented '
    'generation for large language models: A survey '
    '(arXiv:2312.10997). arXiv. https://arxiv.org/abs/2312.10997',

    'He, H., Zhang, H., & Roth, D. (2023). Rethinking with retrieval: '
    'Faithful large language model inference (arXiv:2301.00303). '
    'arXiv. https://arxiv.org/abs/2301.00303',

    'Honovich, O., Aharoni, R., Herzig, J., Taitelbaum, H., '
    'Kukliansy, D., Cohen, V., Scialom, T., & Szpektor, I. (2022). '
    'TRUE: Re-evaluating factual consistency evaluation. In '
    'Proceedings of the 2022 Conference of the North American '
    'Chapter of the Association for Computational Linguistics: '
    'Human Language Technologies (pp. 3905-3920). Association for '
    'Computational Linguistics. '
    'https://aclanthology.org/2022.naacl-main.287/',

    'Huang, L., Yu, W., Ma, W., Zhong, W., Feng, Z., Wang, H., Chen, '
    'Q., Peng, W., Feng, X., Qin, B., & Liu, T. (2025). A survey on '
    'hallucination in large language models: Principles, taxonomy, '
    'challenges, and open questions. ACM Transactions on Information '
    'Systems, 43(2), Article 42. '
    'https://doi.org/10.1145/3703155',

    'Ji, Z., Lee, N., Frieske, R., Yu, T., Su, D., Xu, Y., Ishii, E., '
    'Bang, Y. J., Madotto, A., & Fung, P. (2023). Survey of '
    'hallucination in natural language generation. ACM Computing '
    'Surveys, 55(12), Article 248. '
    'https://doi.org/10.1145/3571730',

    'Kuner, C., Bygrave, L. A., & Docksey, C. (Eds.). (2020). The EU '
    'General Data Protection Regulation (GDPR): A commentary. Oxford '
    'University Press.',

    'Lai, J., Gan, W., Wu, J., Qi, Z., & Yu, P. S. (2024). Large '
    'language models in law: A survey. AI Open, 5, 181-196. '
    'https://doi.org/10.1016/j.aiopen.2024.09.002',

    'Magesh, V., Surani, F., Dahl, M., Suzgun, M., Manning, C. D., & '
    'Ho, D. E. (2024). Hallucination-free? Assessing the reliability '
    'of leading AI legal research tools (arXiv:2405.20362). arXiv. '
    'https://arxiv.org/abs/2405.20362',

    'Malkov, Y. A., & Yashunin, D. A. (2020). Efficient and robust '
    'approximate nearest neighbor search using hierarchical '
    'navigable small world graphs. IEEE Transactions on Pattern '
    'Analysis and Machine Intelligence, 42(4), 824-836. '
    'https://doi.org/10.1109/TPAMI.2018.2889473',

    'Nogueira, R., & Cho, K. (2019). Passage re-ranking with BERT '
    '(arXiv:1901.04086). arXiv. https://arxiv.org/abs/1901.04086',

    'Open Compliance and Ethics Group. (2021). GRC capability model '
    '3.0 (OCEG Red Book). OCEG. '
    'https://www.oceg.org/grc-capability-model-red-book/',

    'Ouyang, L., Wu, J., Jiang, X., Almeida, D., Wainwright, C. L., '
    'Mishkin, P., Zhang, C., Agarwal, S., Slama, K., Ray, A., '
    'Schulman, J., Hilton, J., Kelton, F., Miller, L., Simens, M., '
    'Askell, A., Welinder, P., Christiano, P., Leike, J., & Lowe, R. '
    '(2022). Training language models to follow instructions with '
    'human feedback. In Advances in Neural Information Processing '
    'Systems 35 (pp. 27730-27744). '
    'https://papers.nips.cc/paper_files/paper/2022/hash/'
    'b1efde53be364a73914f58805a001731-Abstract-Conference.html',

    'Thakur, N., Reimers, N., Rücklé, A., Srivastava, A., & '
    'Gurevych, I. (2021). BEIR: A heterogeneous benchmark for '
    'zero-shot evaluation of information retrieval models. In '
    'Proceedings of the 35th Conference on Neural Information '
    'Processing Systems: Datasets and Benchmarks Track. '
    'https://datasets-benchmarks-proceedings.neurips.cc/paper/2021/'
    'hash/65b9eea6e1cc6bb9f0cd2a47751a186f-Abstract-round2.html',

    'Yao, S., Zhao, J., Yu, D., Du, N., Shafran, I., Narasimhan, K., '
    '& Cao, Y. (2023). ReAct: Synergizing reasoning and acting in '
    'language models. In Proceedings of the 11th International '
    'Conference on Learning Representations. '
    'https://openreview.net/forum?id=WE_vluYUL-X',
]


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / '2_1_related_theory.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'2_1_related_theory_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_heading(doc, '2. Background', level=1)
    add_heading(doc, '2.1 Related Theory', level=2)

    add_paragraph(
        doc,
        'This section introduces the four areas of theory the system '
        'relies on: privacy law, information retrieval, language '
        'models, and methods for making AI outputs reliable.',
    )

    add_theory_paragraph(
        doc,
        'Legal and Compliance Foundations.',
        'Modern privacy regulation operates across different '
        'countries, each using its own legal terms and enforcement '
        'methods to manage data processing (Kuner et al., 2020). '
        'Comparative data protection law solves the "comparability '
        'problem" by focusing on whether rules have the same '
        'function, rather than whether they use the same wording. '
        'Legal terminology alignment uses mapping methods to match '
        'similar meanings across jurisdictions. Organisations apply '
        'these rules using Governance, Risk, and Compliance (GRC) '
        'frameworks, which link regulations to internal policies, '
        'controls, audits, and risk management processes (OCEG, '
        '2021). GRC theory focuses on continuous monitoring and '
        'connecting legal requirements to measurable actions inside '
        'an organisation. This theoretical background supports the '
        'system\'s clause-level mapping design, which matches legal '
        'obligations based on meaning and function rather than '
        'exact wording.',
    )

    add_theory_paragraph(
        doc,
        'Information Retrieval and Search Systems.',
        'Information retrieval (IR) theory separates lexical '
        'matching and semantic retrieval. Lexical matching uses '
        'statistical methods such as BM25 to find documents based '
        'on shared words (Thakur et al., 2021). Semantic retrieval '
        'represents text as vectors using transformer-based models '
        'so that meaning, not just words, is compared. Hybrid IR '
        'systems combine both methods to improve accuracy. They '
        'often use fast indexing methods like Hierarchical '
        'Navigable Small World (HNSW) graphs (Malkov & Yashunin, '
        '2020), followed by cross-encoder re-ranking to improve '
        'result quality (Nogueira & Cho, 2019). '
        'Retrieval-Augmented Generation (RAG) improves traditional '
        'retrieval by forcing language models to use retrieved '
        'documents when generating answers, which reduces incorrect '
        'or made-up responses in knowledge-heavy tasks (Y. Gao et '
        'al., 2024). This combined approach supports the '
        'platform\'s document retrieval layer by ensuring that '
        'compliance comparisons are based on real and verifiable '
        'legal texts.',
    )

    add_theory_paragraph(
        doc,
        'Artificial Intelligence and Language Models.',
        'Large language models (LLMs) use transformer-based '
        'architectures that rely on attention mechanisms and are '
        'trained on very large datasets (Bommasani et al., 2021). '
        'They are then improved using instruction tuning so they '
        'can follow complex instructions and produce structured '
        'answers (Ouyang et al., 2022). Instruction tuning also '
        'reduces sensitivity to how prompts are written, allowing '
        'more consistent output for tasks such as compliance '
        'reporting. In legal and compliance settings, LLMs are '
        'useful for understanding meaning in text and extracting '
        'rules or obligations. However, they need careful control '
        'and structure to avoid producing incorrect or unsupported '
        'reasoning (Lai et al., 2024). Agentic workflow theory '
        'describes AI systems that perform tasks in multiple steps, '
        'such as planning, retrieving information, and correcting '
        'their own outputs instead of answering in one step (Yao '
        'et al., 2023). This project uses this agent-based '
        'approach to support tasks like comparing regulations, '
        'mapping policy coverage, and identifying compliance gaps, '
        'while still keeping human review at important stages.',
    )

    add_theory_paragraph(
        doc,
        'Trust, Safety, and Output Reliability.',
        'Hallucination, which is when the model produces fluent '
        'but unsupported or made-up output, is the main failure '
        'mode of LLMs in specialised domains (Huang et al., 2025; '
        'Ji et al., 2023). In legal settings, fabricated citations '
        'have led to real professional consequences (Magesh et al., '
        '2024). Faithfulness means that a generated answer is '
        'supported by its evidence. Grounded generation ensures '
        'this by requiring direct citations to retrieved text '
        'chunks. Verification methods include structural checks '
        'such as citation labels and exact quote matching '
        '(T. Gao et al., 2023), and meaning-based checks using '
        'Natural Language Inference models that decide whether a '
        'statement is supported, not related, or contradicts the '
        'source (He et al., 2023; Honovich et al., 2022). RAG '
        'evaluation tools such as RAGAS measure faithfulness, '
        'answer relevance, and how well the system uses context '
        '(Es et al., 2024).',
    )

    doc.add_page_break()
    add_heading(doc, 'References for §2.1', level=2)
    add_paragraph(
        doc,
        'The following 17 references support the four theory '
        'paragraphs above. Each paragraph uses a disjoint subset of '
        'this list. The full thesis bibliography appears in the '
        'References section at the end of the document.',
    )
    for ref in REFERENCES:
        add_reference(doc, ref)

    doc.save(out_path)
    print(f'Wrote {out_path}')


if __name__ == '__main__':
    main()
