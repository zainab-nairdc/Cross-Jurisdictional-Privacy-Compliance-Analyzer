"""Generate the List of Symbols and List of Abbreviations as Word tables.

Output: thesis_docs/listings_symbols_abbreviations.docx

Both lists are alphabetically sorted and ready to paste into the thesis
front matter.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_run_style(run, *, bold=False, italic=False, color=NAVY, size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 1', 2: 'Heading 2'}
    size_map = {1: 16, 2: 14}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_table(doc, headers, rows, widths_cm, body_size=11):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for col_idx, w in enumerate(widths_cm):
        for cell in table.columns[col_idx].cells:
            cell.width = Cm(w)

    header_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = h
        shade_cell(cell, '002583')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            set_run_style(run, bold=True, color=WHITE, size=11)

    for r_idx, row_values in enumerate(rows, start=1):
        row = table.rows[r_idx]
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_style(
                        run,
                        bold=(c_idx == 0),
                        color=DARK_GRAY,
                        size=body_size,
                    )


SYMBOLS = [
    ('~', 'Approximately (used with corpus and document counts)'),
    ('≥', 'Greater than or equal to (used with performance targets)'),
    ('✔', 'Engaged or applies (used in requirements tables under '
     'Data, Process, Communication)'),
]


ABBREVIATIONS = [
    ('AI', 'Artificial Intelligence'),
    ('API', 'Application Programming Interface'),
    ('ASGI', 'Asynchronous Server Gateway Interface'),
    ('BBK', 'Bank of Bahrain and Kuwait'),
    ('BGE', 'BAAI General Embedding'),
    ('BM25', 'Best Matching 25'),
    ('BP', 'Bahrain Polytechnic'),
    ('CI', 'Continuous Integration'),
    ('CJPCA', 'Cross-Jurisdictional Privacy Compliance Analyzer'),
    ('CLI', 'Command Line Interface'),
    ('CON', 'Constraint'),
    ('CPU', 'Central Processing Unit'),
    ('CSP', 'Content Security Policy'),
    ('CSRF', 'Cross-Site Request Forgery'),
    ('CSS', 'Cascading Style Sheets'),
    ('DOCX', 'Microsoft Word Open XML Document'),
    ('DPDP', 'Digital Personal Data Protection (Act, India)'),
    ('FTS5', 'Full-Text Search version 5'),
    ('GDPR', 'General Data Protection Regulation'),
    ('GPU', 'Graphics Processing Unit'),
    ('GRC', 'Governance, Risk, and Compliance'),
    ('HNSW', 'Hierarchical Navigable Small World'),
    ('HSTS', 'HTTP Strict Transport Security'),
    ('HTML', 'HyperText Markup Language'),
    ('HTTP', 'HyperText Transfer Protocol'),
    ('HTTPS', 'HTTP Secure'),
    ('ICT', 'Information and Communication Technology'),
    ('IDE', 'Integrated Development Environment'),
    ('IEC', 'International Electrotechnical Commission'),
    ('IP', 'Internet Protocol'),
    ('IR', 'Information Retrieval'),
    ('ISO', 'International Organization for Standardization'),
    ('JSON', 'JavaScript Object Notation'),
    ('LLM', 'Large Language Model'),
    ('MFA', 'Multi-Factor Authentication'),
    ('MoSCoW', 'Must, Should, Could, Won\'t have (prioritisation method)'),
    ('NAIRDC',
     'Nasser Artificial Intelligence Research and Development Centre'),
    ('NFR', 'Non-Functional Requirement'),
    ('NIST', 'National Institute of Standards and Technology'),
    ('NLI', 'Natural Language Inference'),
    ('NLP', 'Natural Language Processing'),
    ('ONNX', 'Open Neural Network Exchange'),
    ('ORM', 'Object-Relational Mapping'),
    ('OTP', 'One-Time Password'),
    ('OWASP', 'Open Web Application Security Project'),
    ('PDF', 'Portable Document Format'),
    ('PDPL', 'Personal Data Protection Law (Bahrain)'),
    ('RAG', 'Retrieval-Augmented Generation'),
    ('RAGAS', 'Retrieval-Augmented Generation Assessment System'),
    ('RBAC', 'Role-Based Access Control'),
    ('RBI', 'Reserve Bank of India'),
    ('Req', 'Requirement'),
    ('SaaS', 'Software as a Service'),
    ('SHA', 'Secure Hash Algorithm'),
    ('SOP', 'Standard Operating Procedure'),
    ('SQL', 'Structured Query Language'),
    ('SVN', 'Subversion'),
    ('TLS', 'Transport Layer Security'),
    ('TOC', 'Table of Contents'),
    ('TOF', 'Table of Figures'),
    ('TOT', 'Table of Tables'),
    ('TOTP', 'Time-Based One-Time Password'),
    ('UML', 'Unified Modeling Language'),
    ('URL', 'Uniform Resource Locator'),
    ('USR', 'User Story Requirement'),
    ('VS Code', 'Visual Studio Code'),
    ('WAL', 'Write-Ahead Logging'),
    ('XLSX', 'Microsoft Excel Open XML Spreadsheet'),
]


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / 'listings_symbols_abbreviations.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = (out_dir /
                             f'listings_symbols_abbreviations_v{n}.docx')
                if not candidate.exists():
                    out_path = candidate
                    break

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_heading(doc, 'List of Symbols', level=1)
    add_paragraph(doc, 'In alphabetical order.', italic=True,
                  space_after=10)
    add_table(
        doc,
        headers=['Symbol', 'Meaning'],
        rows=SYMBOLS,
        widths_cm=[2.5, 13.5],
        body_size=11,
    )

    doc.add_paragraph()

    add_heading(doc, 'List of Abbreviations', level=1)
    add_paragraph(doc, 'In alphabetical order.', italic=True,
                  space_after=10)
    add_table(
        doc,
        headers=['Abbreviation', 'Full Form'],
        rows=ABBREVIATIONS,
        widths_cm=[3.5, 12.5],
        body_size=11,
    )

    doc.save(out_path)
    print(f'Wrote {out_path}')


if __name__ == '__main__':
    main()
