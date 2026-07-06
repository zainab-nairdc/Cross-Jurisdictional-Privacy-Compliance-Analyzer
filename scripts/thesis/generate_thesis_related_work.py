"""
Generate the "Related Work" thesis section as a .docx file.

Produces: thesis_related_work.docx in the project root.

Body word budget: ~400 words (excluding references list).
Format mirrors the user's previous thesis Related Work section: short
sub-headings, prose paragraphs, in-text (Author, Year) citations, and a
References list at the end.
"""

from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = PROJECT_ROOT / "thesis_related_work.docx"

SECTION_HEADING = "Related Work"

OPENING_PARAGRAPH = (
    "The Cross-Jurisdictional Privacy Compliance Analyzer draws on three "
    "converging research streams: retrieval-augmented generation in specialised "
    "domains, hallucination mitigation in legal artificial intelligence, and "
    "comparative privacy law. The literature within each stream is mature in "
    "isolation but, as discussed below, has not been integrated for the "
    "specific problem of cross-jurisdictional regulatory alignment."
)

SUBSECTIONS = [
    (
        "Retrieval-Augmented Generation in Specialised Domains",
        "Lewis et al. (2020) introduced retrieval-augmented generation (RAG) as a "
        "method for grounding generative language models in retrieved evidence, "
        "sharply reducing hallucination on knowledge-intensive tasks. Subsequent "
        "work has refined the retrieval component: Karpukhin et al. (2020) "
        "demonstrated that dense passage retrievers outperform classical BM25 "
        "(Robertson & Zaragoza, 2009) on open-domain questions, while hybrid "
        "systems using Reciprocal Rank Fusion (Cormack et al., 2009) consistently "
        "lead modern benchmarks. Gao et al. (2024) survey current RAG architectures "
        "and identify multi-step retrieval and domain-specific tuning as essential "
        "for compositional, regulation-style queries. Domain adaptation is "
        "well-established in legal NLP, where corpus-specific encoders such as "
        "LEGAL-BERT (Chalkidis et al., 2020) yield material gains over "
        "general-purpose models, although performance remains bounded by "
        "passage-level retrieval and single-jurisdiction training corpora."
    ),
    (
        "Hallucination and Evaluation in Legal AI",
        "Hallucination in language-model output is a documented and persistent "
        "failure mode (Ji et al., 2023) with particular consequences in legal "
        "contexts. Dahl et al. (2024) measured legal-hallucination rates of 69 per "
        "cent for ChatGPT 3.5 and 88 per cent for Llama 2 on questions about US "
        "case law, demonstrating that fluency does not imply faithfulness. Even "
        "commercial RAG-based legal-research products are not exempt: Magesh et "
        "al. (2024) report that Lexis+ AI hallucinates on 17 per cent of grounded "
        "queries and Westlaw AI on 33 per cent. Automated evaluation frameworks "
        "such as RAGAS (Es et al., 2024) score faithfulness and answer relevancy "
        "without requiring labelled ground truth, providing the methodological "
        "scaffolding that the present project adopts to bound residual "
        "non-determinism."
    ),
    (
        "Comparative Privacy Law and Compliance Automation",
        "Greenleaf (2023) catalogues the global proliferation of privacy regimes, "
        "now exceeding 162 national laws, noting that the rate of new legislation "
        "has outpaced doctrinal analytical capacity, particularly across emerging "
        "Asian and Gulf jurisdictions. Computational privacy work has historically "
        "focused on consumer-facing privacy-policy parsing (Wilson et al., 2016) "
        "rather than regulation-to-policy mapping or cross-jurisdictional "
        "alignment, leaving a methodological gap for systems that operate at the "
        "level of statutory text across multiple regimes."
    ),
    (
        "Research Gap",
        "Existing legal-AI systems are typically single-jurisdiction and "
        "clause-aligned, while comparative privacy scholarship remains largely "
        "doctrinal. No published system combines retrieval-grounded reasoning, "
        "taxonomy-based topic alignment, and citation-verified output across "
        "asymmetric regulatory regimes such as Bahrain’s PDPL, India’s DPDPA, and "
        "Kuwait’s DPPR."
    ),
]


REFERENCES = [
    "Chalkidis, I., Fergadiotis, M., Malakasiotis, P., Aletras, N., & Androutsopoulos, I. (2020). LEGAL-BERT: The Muppets straight out of Law School. Findings of the Association for Computational Linguistics: EMNLP 2020, 2898–2904.",
    "Cormack, G. V., Clarke, C. L. A., & Büttcher, S. (2009). Reciprocal rank fusion outperforms Condorcet and individual rank learning methods. Proceedings of the 32nd International ACM SIGIR Conference on Research and Development in Information Retrieval (SIGIR '09), 758–759.",
    "Dahl, M., Magesh, V., Suzgun, M., & Ho, D. E. (2024). Large legal fictions: Profiling legal hallucinations in large language models. Journal of Legal Analysis, 16(1), 64–93.",
    "Es, S., James, J., Espinosa-Anke, L., & Schockaert, S. (2024). RAGAs: Automated evaluation of retrieval augmented generation. Proceedings of the 18th Conference of the European Chapter of the Association for Computational Linguistics: System Demonstrations (EACL 2024), 150–158.",
    "Gao, Y., Xiong, Y., Gao, X., Jia, K., Pan, J., Bi, Y., Dai, Y., Sun, J., Wang, M., & Wang, H. (2024). Retrieval-augmented generation for large language models: A survey. arXiv preprint arXiv:2312.10997.",
    "Greenleaf, G. (2023). Global data privacy laws 2023: 162 national laws and 20 bills. Privacy Laws & Business International Report, 181, 2–4.",
    "Ji, Z., Lee, N., Frieske, R., Yu, T., Su, D., Xu, Y., Ishii, E., Bang, Y. J., Madotto, A., & Fung, P. (2023). Survey of hallucination in natural language generation. ACM Computing Surveys, 55(12), Article 248.",
    "Karpukhin, V., Oğuz, B., Min, S., Lewis, P., Wu, L., Edunov, S., Chen, D., & Yih, W. (2020). Dense passage retrieval for open-domain question answering. Proceedings of the 2020 Conference on Empirical Methods in Natural Language Processing (EMNLP), 6769–6781.",
    "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., Küttler, H., Lewis, M., Yih, W., Rocktäschel, T., Riedel, S., & Kiela, D. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. Advances in Neural Information Processing Systems 33 (NeurIPS 2020), 9459–9474.",
    "Magesh, V., Surani, F., Dahl, M., Suzgun, M., Manning, C. D., & Ho, D. E. (2024). Hallucination-free? Assessing the reliability of leading AI legal research tools. Stanford RegLab preprint (forthcoming, Journal of Empirical Legal Studies).",
    "Robertson, S., & Zaragoza, H. (2009). The probabilistic relevance framework: BM25 and beyond. Foundations and Trends in Information Retrieval, 3(4), 333–389.",
    "Wilson, S., Schaub, F., Dara, A. A., Liu, F., Cherivirala, S., Leon, P. G., Andersen, M. S., Zimmeck, S., Sathyendra, K. M., Russell, N. C., Norton, T. B., Hovy, E., Reidenberg, J., & Sadeh, N. (2016). The creation and analysis of a website privacy policy corpus. Proceedings of the 54th Annual Meeting of the Association for Computational Linguistics (ACL 2016), 1330–1340.",
]


def _count_body_words() -> int:
    total = len(OPENING_PARAGRAPH.split())
    for _, body in SUBSECTIONS:
        total += len(body.split())
    return total


def build_document() -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(SECTION_HEADING, level=1)

    opening = doc.add_paragraph(OPENING_PARAGRAPH)
    opening.paragraph_format.space_after = Pt(10)

    for heading, body in SUBSECTIONS:
        doc.add_heading(heading, level=2)
        para = doc.add_paragraph(body)
        para.paragraph_format.space_after = Pt(10)

    doc.add_heading("References", level=2)
    for entry in REFERENCES:
        ref_para = doc.add_paragraph(entry)
        ref_para.paragraph_format.left_indent = Cm(1.0)
        ref_para.paragraph_format.first_line_indent = Cm(-1.0)
        ref_para.paragraph_format.space_after = Pt(6)

    return doc


def main() -> None:
    word_count = _count_body_words()
    print(f"Body word count (excluding references): {word_count}")
    doc = build_document()
    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
