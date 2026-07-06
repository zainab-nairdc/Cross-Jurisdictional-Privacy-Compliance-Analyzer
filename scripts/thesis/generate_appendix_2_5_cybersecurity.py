"""Generate Appendix 2.5 Cybersecurity Supplements as a Word document.

Output: thesis_docs/appendix_2_5_cybersecurity.docx

Contents:
  A.2.5 Cybersecurity Supplements
    A.2.5.1 Risk Assessment Register     (Table A.2.5.1)
    A.2.5.2 Security Policies            (Table A.2.5.2)
    A.2.5.3 Logs and Detection Flow      (Figure A.2.5.1)

Style follows the existing thesis appendices (navy headings, 2 cm margins).
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)


# Table A.2.5.1 — Risk Assessment Register
# Rows from §5.3, condensed for the appendix layout
# (id, risk, vector, likelihood, impact, mitigation, residual)
RISK_ROWS = [
    ('R-01', 'Credential compromise',
     'Brute force, credential stuffing', 'Medium', 'High',
     'MFA + password validators + django-axes lockout', 'Low'),
    ('R-02', 'Session theft',
     'XSS or network MITM', 'Low', 'High',
     'HttpOnly + Secure + SameSite cookies, HSTS, CSP', 'Low'),
    ('R-03', 'CSRF on a mutating action',
     'Phished link, forged POST', 'Low', 'High',
     'CsrfViewMiddleware + token per form', 'Low'),
    ('R-04', 'Stored XSS via document content',
     'Malicious file payload', 'Low', 'High',
     'Template autoescape; CSP without unsafe-inline scripts', 'Low'),
    ('R-05', 'SQL injection',
     'Crafted query', 'Very low', 'High',
     'Django ORM with parameter binding throughout', 'Negligible'),
    ('R-06', 'File-upload malware',
     'Malicious PDF or DOCX', 'Low', 'High',
     'Suffix allowlist, SHA-256 dedup, no script execution path', 'Low'),
    ('R-07', 'Prompt injection via corpus',
     'Attacker-controlled chunk text', 'Medium', 'Medium',
     'Tier-A regex + Tier-B LLM judge + verifier auto-drop', 'Low'),
    ('R-08', 'Hallucinated citation',
     'Statistical LLM behaviour', 'Medium', 'High',
     'Verbatim grounding, NLI score, auto-drop above threshold', 'Low'),
    ('R-09', 'API-key leakage',
     'Logs or error pages', 'Low', 'Critical',
     'Env-vars only, DEBUG=False, sanitised audit', 'Low'),
    ('R-10', 'LLM provider outage',
     'External vendor', 'Low', 'Medium',
     'OutputFixingParser + SafeFallback empty report', 'Medium'),
    ('R-11', 'Insider exfiltration of corpus',
     'Login + scope abuse', 'Low', 'High',
     'RBAC, per-row scope, audit-log of document views', 'Low'),
    ('R-12', 'Supply-chain pip dependency',
     'Backdoor in transitive dep', 'Low', 'Critical',
     'Pinned versions, periodic SBOM and CVE review', 'Medium'),
    ('R-13', 'Supply-chain CDN',
     'Backdoor in static asset', 'Low', 'High',
     'CSP allowlist; SRI hashes recommended', 'Medium'),
    ('R-14', 'Audit-log tampering by admin',
     'Direct DB write by insider', 'Low', 'Critical',
     'App-layer append-only; admin actions themselves audited', 'Medium'),
    ('R-15', 'DoS via large upload',
     'Oversized payload', 'Low', 'Low',
     'DATA_UPLOAD_MAX_MEMORY_SIZE cap; proxy body limit', 'Low'),
    ('R-16', 'Geographic policy violation',
     'Login from disallowed country', 'Low', 'Medium',
     'GeoFenceMiddleware with MaxMind allowlist', 'Low'),
    ('R-17', 'Idle-session takeover',
     'Unlocked browser', 'Low', 'Medium',
     'IdleSessionTimeoutMiddleware (1 800 s)', 'Low'),
    ('R-18', 'Forced password change skipped',
     'Misconfiguration', 'Low', 'Medium',
     'ForcePasswordChangeMiddleware on must_change', 'Low'),
    ('R-19', 'MFA not enforced',
     'Misconfiguration', 'Low', 'High',
     'ForceMFAEnrollmentMiddleware', 'Low'),
    ('R-20', 'Reasoning regression on model upgrade',
     'Vendor model change', 'Medium', 'Medium',
     'Eval-snapshot regression test (recommended CI gate)', 'Medium'),
]


# Table A.2.5.2 — Security Policies
POLICY_ROWS = [
    ('Password policy',
     'Min 12 chars, complexity (upper+lower+digit+special), '
     'similarity ≤ 0.5 to username, common-password reject, '
     'breached-list check.'),
    ('Session policy',
     'SESSION_COOKIE_AGE = 28 800 s; Secure + HttpOnly + SameSite; '
     'idle timeout 1 800 s; logout on browser close.'),
    ('MFA policy',
     'TOTP required for every authenticated user; backup tokens via '
     'django_otp.plugins.otp_static; production REQUIRE_MFA = True.'),
    ('RBAC policy',
     'Three roles (analyst, reviewer, admin); role grant by admin only; '
     'role change is itself an audited action.'),
    ('Audit retention policy',
     'AuditLog rows are append-only at the application layer; retained '
     'indefinitely; rotated per-job log files retained operationally '
     '(suggested 30 days).'),
    ('Geofence policy',
     'Default allowlist: Bahrain, India, Kuwait, United Arab Emirates; '
     'fail-open on lookup error to avoid lock-out from BBK network.'),
    ('File upload policy',
     'Suffix allowlist PDF + DOCX; size cap '
     'DATA_UPLOAD_MAX_MEMORY_SIZE; SHA-256 content hash for dedup; '
     'no OCR script execution path.'),
    ('Dependency review policy',
     'Pinned requirements.txt; pip-audit recommended monthly; '
     'Django security advisories tracked; HuggingFace model versions '
     'pinned and reviewed against the RAGAS eval suite before upgrade.'),
]


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_run_style(run, *, bold=False, italic=False, color=NAVY, size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 2', 2: 'Heading 3', 3: 'Heading 4'}
    size_map = {1: 14, 2: 12, 3: 11}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_caption(doc, label: str, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    r1 = p.add_run(label + ' ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run(caption)
    set_run_style(r2, italic=True, color=MUTED, size=10)


def add_image(doc, image_path: Path, width_cm: float = 16.0):
    if not image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'[MISSING IMAGE: {image_path.name}]')
        set_run_style(run, bold=True, italic=True, color=MUTED, size=10)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def add_table(doc, headers, rows, widths_cm, body_size=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for col_idx, w in enumerate(widths_cm):
        for cell in table.columns[col_idx].cells:
            cell.width = Cm(w)

    header_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = h
        shade_cell(cell, '002583')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            set_run_style(run, bold=True, color=WHITE, size=10)

    for r_idx, row_values in enumerate(rows, start=1):
        row = table.rows[r_idx]
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_style(
                        run,
                        bold=(c_idx == 0),
                        color=DARK_GRAY,
                        size=body_size,
                    )


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / 'appendix_2_5_cybersecurity.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'appendix_2_5_cybersecurity_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    diagrams = base / 'diagrams'
    fig_a251 = diagrams / 'Figure_A_2_5_1_LogsDetection.drawio.png'

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_heading(doc, 'Appendix 2.5 Cybersecurity Supplements', level=1)

    add_paragraph(
        doc,
        'This appendix supports §3.2.5 with the risk assessment register, '
        'the in-system security policies, and the logs-and-detection flow. '
        'Full threat modelling and secure development practices are in '
        'Chapter 5.',
    )

    # A.2.5.1 Risk Assessment Register
    add_heading(doc, 'A.2.5.1 Risk Assessment Register', level=2)
    add_caption(doc, 'Table A.2.5.1.', 'Risk Assessment Register')
    add_table(
        doc,
        ['ID', 'Risk', 'Vector', 'Likelihood', 'Impact', 'Mitigation',
         'Residual'],
        RISK_ROWS,
        widths_cm=[1.2, 2.8, 2.8, 1.6, 1.4, 4.2, 1.4],
        body_size=8,
    )

    # A.2.5.2 Security Policies
    add_heading(doc, 'A.2.5.2 Security Policies', level=2)
    add_caption(doc, 'Table A.2.5.2.', 'Security Policies')
    add_table(
        doc,
        ['Policy', 'Statement'],
        POLICY_ROWS,
        widths_cm=[4.0, 12.0],
    )

    # A.2.5.3 Logs and Detection Flow
    add_heading(doc, 'A.2.5.3 Logs and Detection Flow', level=2)
    add_paragraph(
        doc,
        'Figure A.2.5.1 traces an in-system action from the user, through '
        'the middleware stack, into the audit signal, and out to the '
        'inspection sinks. Lockout, geo-block, and RBAC denials follow the '
        'short path on the bottom of the figure and never reach the audit '
        'persistence path; they are recorded by the originating middleware.',
    )
    add_image(doc, fig_a251, width_cm=16.0)
    add_caption(doc, 'Figure A.2.5.1.', 'Logs and Detection Flow')

    doc.save(out_path)
    print(f'Wrote: {out_path}')


if __name__ == '__main__':
    main()
