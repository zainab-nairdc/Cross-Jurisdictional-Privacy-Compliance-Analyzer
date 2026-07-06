"""Generate Appendix 2.1 Requirements Traceability Matrix as a Word document.

Output: thesis_docs/appendix_2_1_traceability.docx

Contents:
  1. Appendix 2.1 heading
  2. Short lead-in paragraph
  3. Table A2.1.1 — Functional Requirements (20 rows)
  4. Table A2.1.2 — Non-Functional Requirements (27 rows)
  5. Table A2.1.3 — Compliance Requirements (9 rows)

Each table has four columns: ID, Requirement, Design Element, Source.
Style mirrors the Stage 1 + Stage 2 generators (navy headings, 2 cm margins,
Table Grid). Pulls requirement text from thesis_docs/02_requirements.md
indirectly — the rows are hardcoded here for stability.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)


# ── Functional Requirements (20 rows) ─────────────────────────────────────
FR_ROWS = [
    ('FR-01', 'Document ingestion (PDF/DOCX)',
     '§3.2.1.5; Figure 4 (Admin); Figure A2.3d',
     'NAIRDC brief'),
    ('FR-02', 'Hybrid retrieval',
     '§3.2.1 Table 7 rows 1–3; Figure 6; Figure A2.3e; §3.2.3 (forthcoming)',
     'Prototype evaluation'),
    ('FR-03', 'Cross-jurisdictional comparison',
     'Figure 7 (ComparisonRun, ComparisonResult); Figure A2.5c; §3.2.3 (forthcoming)',
     'NAIRDC brief'),
    ('FR-04', 'Policy-to-regulation mapping',
     'Figure 7 (MappingAnalysis, ObligationMapping); Figure A2.5a–b; Figure A2.3a; §3.2.3 (forthcoming)',
     'BBK consultation'),
    ('FR-05', 'Gap analysis',
     'Figure 7 (Gap class); §3.2.3 (forthcoming)',
     'BBK consultation'),
    ('FR-06', 'Auto-routed mapping',
     '§3.2.1 Table 7 row 8 (chunk_tags taxonomy); §3.2.3 (forthcoming)',
     'Prototype evaluation'),
    ('FR-07', 'Citation verification',
     '§3.2.1 Table 7 row 6; Figure 6; Figure A2.3f; §3.2.3 (forthcoming)',
     'Prototype evaluation'),
    ('FR-08', 'Copilot Q&A',
     'Figure 4 (Ask Copilot); Figure A2.3b; §3.2.3 (forthcoming)',
     'BBK consultation'),
    ('FR-09', 'User authentication (MFA)',
     'Figure 5 (Django + 2FA); §3.2.4 (forthcoming)',
     'Financial-sector baseline'),
    ('FR-10', 'Role-based access control',
     '§3.2.1 (UserProfile.role); Figure 4 (three actors); Figure 7 (Role enum)',
     'NAIRDC brief'),
    ('FR-11', 'Reviewer accept / reject / modify',
     'Figure 4 (Reviewer); Figure A2.3c (cascade); Figure A2.5b',
     'BBK consultation'),
    ('FR-12', 'Audit log',
     '§3.2.1 Table 7 rows 9–10; Figure 7 (AuditLog); §3.2.1.5 (audit-hash chain)',
     'STRIDE; PDPL Art. 21'),
    ('FR-13', 'Analytics dashboard',
     'Figure 4 (View analytics dashboard)',
     'BBK consultation'),
    ('FR-14', 'Library workspace',
     'Figure 4 (Browse library)',
     'BBK consultation'),
    ('FR-15', 'Term Dictionary',
     'Figure 4 (Browse library); §3.2.1 Table 7 row 8',
     'BBK consultation'),
    ('FR-16', 'Async job runner',
     'Figure 5 (Background workers); Figure 7 (job models); Figure A2.5a',
     'Operational constraint'),
    ('FR-17', 'Export reports',
     'Figure 4 (Download approved report); §3.2.1.5 (audit-hash chain)',
     'BBK consultation'),
    ('FR-18', 'Document upload',
     'Figure 4 (Admin → Upload documents); Figure A2.3d',
     'NAIRDC brief'),
    ('FR-19', 'Notifications',
     'Figure 7 (job-status fields drive UI badges)',
     'BBK consultation'),
    ('FR-20', 'History',
     'Figure 4 (View audit log); Figure 7 (AuditLog)',
     'BBK consultation'),
]


# ── Non-Functional Requirements (27 rows) ─────────────────────────────────
NFR_ROWS = [
    ('NFR-01', 'Hybrid retrieval ≤ 1.5 s',
     '§3.2.1 Table 7 rows 1–3 (in-process indexes); Figure A2.3e',
     'Prototype evaluation'),
    ('NFR-02', 'Comparison ≤ 3 min',
     '§3.2.3 (forthcoming)',
     'Prototype evaluation'),
    ('NFR-03', 'Mapping ≤ 2 min',
     '§3.2.3 (forthcoming)',
     'Prototype evaluation'),
    ('NFR-04', 'Gap analysis ≤ 5 min',
     '§3.2.3 (forthcoming)',
     'Prototype evaluation'),
    ('NFR-05', 'Corpus of up to 3,000 chunks',
     '§3.2.1 Table 7 row 1 (ChromaDB scales to 50k)',
     'NAIRDC brief'),
    ('NFR-06', '100 concurrent users',
     'Figure 5 (single-host shape); §3.2.4 (forthcoming)',
     'NAIRDC brief'),
    ('NFR-07', '99% uptime in working hours',
     'Figure 5; §3.2.4 (forthcoming)',
     'Operational constraint'),
    ('NFR-08', 'No data loss on power failure',
     '§3.2.1 Table 7 (SQLite WAL); Figure 5 (persistent stores)',
     'Operational constraint'),
    ('NFR-09', 'Auto-correction ≥ 80% of citation drift',
     '§3.2.1 Table 7 row 6; Figure A2.3f',
     'Prototype evaluation'),
    ('NFR-10', 'Auth on all non-public routes',
     'Figure 5 (Django + 2FA); §3.2.4 (forthcoming)',
     'STRIDE'),
    ('NFR-11', 'Password ≥ 12 chars with complexity checks',
     '§3.2.4 (forthcoming)',
     'Financial-sector baseline'),
    ('NFR-12', 'Brute-force lockout',
     '§3.2.4 (forthcoming)',
     'STRIDE'),
    ('NFR-13', 'Outbound traffic restricted',
     'Figure 5 (trust boundary); §3.2.4 (forthcoming)',
     'STRIDE'),
    ('NFR-14', 'New jurisdictions via config only',
     '§3.2.1 (Document.jurisdiction enum); §3.2.4 (forthcoming)',
     'NAIRDC brief'),
    ('NFR-15', 'Type-checked module boundaries',
     '§3.2.1 Table 7 row 7 (Pydantic v2)',
     'Operational constraint'),
    ('NFR-16', 'LLM provider switchable',
     '§3.2.4 (forthcoming)',
     'NAIRDC brief'),
    ('NFR-17', 'Per-job structured log',
     'Figure 7 (IngestionJob.log_entries)',
     'Operational constraint'),
    ('NFR-18', 'Progress visible within 2 s',
     'Figure 7 (IngestionJob.progress_pct); Figure A2.5a',
     'Operational constraint'),
    ('NFR-19', 'Outputs include verbatim quotes',
     '§3.2.1.5; Figure 7 (ObligationMapping.evidence_text)',
     'NAIRDC brief; PDPL audit clauses'),
    ('NFR-20', 'Audit log append-only',
     '§3.2.1.5; Figure 7 (AuditLog)',
     'STRIDE; PDPL Art. 21'),
    ('NFR-21', 'Internal policies stay on host',
     'Figure 5 (trust boundary)',
     'BBK consultation'),
    ('NFR-22', '≤ 4 clicks per workflow',
     'Figure 4 (use case flow)',
     'BBK consultation'),
    ('NFR-23', 'WCAG-AA accessibility',
     '§3.2.4 (forthcoming)',
     'Operational constraint'),
    ('NFR-24', 'Models cached locally on disk',
     'Figure 5 (host-local cache)',
     'Operational constraint'),
    ('NFR-25', 'Recoverable in 30 min on fresh host',
     '§3.2.4 (forthcoming)',
     'Operational constraint'),
    ('NFR-26', 'Failed LLM retried then fallback',
     '§3.2.1 Table 7 row 6; Figure A2.3f',
     'Prototype evaluation'),
    ('NFR-27', 'Workspace pages render ≤ 500 ms',
     'Figure 5 (server-side render); §3.2.4 (forthcoming)',
     'Operational constraint'),
]


# ── Compliance Requirements (9 rows) ──────────────────────────────────────
CR_ROWS = [
    ('CR-01', 'Maintain auditable processing records',
     '§3.2.1.5 (audit-hash chain); Figure 7 (AuditLog)',
     'PDPL Art. 21; DPDPA s. 8(5); DPPR Art. 7'),
    ('CR-02', 'Identify legal basis per finding',
     'Figure 7 (ObligationMapping.regulation_evidence)',
     'PDPL Art. 4; DPDPA s. 6; DPPR Art. 4'),
    ('CR-03', 'Cross-border transfer rules per jurisdiction',
     '§3.2.1 Table 7 row 8 (Cross_Border_Transfer topic)',
     'PDPA Order 42/2022; DPDPA s. 16; DPPR Art. 19'),
    ('CR-04', 'Breach-notification timelines per jurisdiction',
     '§3.2.1 Table 7 row 8 (Breach_Notification_Timeline topic)',
     'PDPL Art. 11; DPDPA s. 8(6); DPPR Art. 16'),
    ('CR-05', 'DPO appointment and duties',
     '§3.2.1 Table 7 row 8 (DPO topic)',
     'PDPA Order 46/2022; DPDPA s. 10; DPPR Art. 9'),
    ('CR-06', 'Sensitive-data special protections',
     '§3.2.1 Table 7 row 8 (Sensitive_Data topic)',
     'PDPA Order 45/2022; DPDPA s. 2(t)(u); DPPR Art. 11'),
    ('CR-07', 'Data-subject rights coverage',
     '§3.2.1 Table 7 row 8 (Data_Subject_Rights group)',
     'PDPL Art. 18–20; DPDPA s. 11–13; DPPR Art. 14–17'),
    ('CR-08', 'Evidentiary integrity of reports',
     '§3.2.1.5 (audit-hash chain); Figure 7 (lifecycle FSM); Figure A2.5b',
     'Audit clauses across all three regimes'),
    ('CR-09', 'Disclose AI involvement in outputs',
     'Figure 4 (AI Drafted / Edited / Human Written annotations)',
     'Responsible-AI best practice; DPDPA accountability principles'),
]


# ── Style helpers ─────────────────────────────────────────────────────────

def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_run_style(run, *, bold=False, italic=False, color=NAVY, size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 2', 2: 'Heading 3'}
    size_map = {1: 14, 2: 12}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_table_caption(doc, label: str, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(label + ' ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run(caption)
    set_run_style(r2, italic=True, color=MUTED, size=10)


def add_traceability_table(doc, rows):
    """Render a 4-column traceability table.

    Columns: ID | Requirement | Design Element | Source.
    """
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [Cm(1.6), Cm(4.6), Cm(6.4), Cm(4.0)]
    for col_idx, w in enumerate(widths):
        for cell in table.columns[col_idx].cells:
            cell.width = w

    headers = ['ID', 'Requirement', 'Design Element', 'Source']
    header_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = h
        shade_cell(cell, '002583')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            set_run_style(run, bold=True, color=WHITE, size=10)

    for r_idx, (rid, req, design, source) in enumerate(rows, start=1):
        row = table.rows[r_idx]
        row.cells[0].text = rid
        row.cells[1].text = req
        row.cells[2].text = design
        row.cells[3].text = source
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_style(
                        run,
                        bold=(c_idx == 0),
                        color=DARK_GRAY,
                        size=9,
                    )


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / 'appendix_2_1_traceability.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'appendix_2_1_traceability_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_heading(doc, 'Appendix 2.1 — Requirements Traceability Matrix', level=1)

    add_paragraph(
        doc,
        'This appendix maps each requirement from §3.1 to the section, '
        'figure, or table in §3.2 that satisfies it, and notes the source '
        'the requirement was derived from. Functional, non-functional, '
        'and compliance requirements are tabulated separately. References '
        'marked "forthcoming" point to sections of §3.2 not yet covered '
        'in this draft; they will be filled in once those sections are '
        'finalised.',
    )

    # ── Table A2.1.1 — Functional ─────────────────────────────────────────
    add_table_caption(
        doc,
        'Table A2.1.1.',
        'Functional Requirements Traceability.',
    )
    add_traceability_table(doc, FR_ROWS)

    # ── Table A2.1.2 — Non-Functional ─────────────────────────────────────
    add_table_caption(
        doc,
        'Table A2.1.2.',
        'Non-Functional Requirements Traceability.',
    )
    add_traceability_table(doc, NFR_ROWS)

    # ── Table A2.1.3 — Compliance ─────────────────────────────────────────
    add_table_caption(
        doc,
        'Table A2.1.3.',
        'Compliance Requirements Traceability.',
    )
    add_traceability_table(doc, CR_ROWS)

    doc.save(out_path)
    print(f'Wrote: {out_path}')
    print(f'  Functional rows:     {len(FR_ROWS)}')
    print(f'  Non-functional rows: {len(NFR_ROWS)}')
    print(f'  Compliance rows:     {len(CR_ROWS)}')


if __name__ == '__main__':
    main()
