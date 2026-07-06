"""Generate the Non-Functional Requirements Word document for the thesis.

Output: thesis_non_functional_requirements.docx in the project root.

Same layout as the FR doc: one continuous table with four columns
(ID, Priority, Source, Requirement Description). Category names appear as
merged-cell section headers between requirement groups.
Priority: H (must-have), M (should-have), L (could-have) — MoSCoW.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


CATEGORIES = [
    (
        'Performance Efficiency',
        [
            ('H', 'Technical feasibility study', 'The system shall return retrieval results within 2 seconds for typical queries on a warm cache.'),
            ('M', 'Technical feasibility study', 'The system shall complete a single-topic regulation comparison within 60 seconds for a typical regulation pair.'),
            ('M', 'Technical feasibility study', 'The system shall complete a single-topic policy mapping within 90 seconds for a typical policy document.'),
            ('M', 'Discussion',                  'The system shall update background-job progress at a polling interval of no more than 2 seconds.'),
            ('L', 'Technical feasibility study', 'The system shall cache navigation badge counts for at least 30 seconds to minimise repeated database queries.'),
        ],
    ),
    (
        'Scalability',
        [
            ('H', 'Existing system analysis',   'The system shall support concurrent use by at least 100 active users at the BBK deployment scale.'),
            ('H', 'Technical feasibility study','The system shall support a regulatory corpus of at least 50,000 indexed chunks without degradation of retrieval response time.'),
            ('L', 'Discussion',                 'The system shall scale vertically on a single host; horizontal scaling is acknowledged as a future-work constraint.'),
        ],
    ),
    (
        'Reliability & Availability',
        [
            ('H', 'Research paper',             'The system shall execute reasoning workflows as a bounded iterative correction process.'),
            ('M', 'Research paper',             'The system shall produce a deterministic fallback response when reasoning workflows cannot produce a verified answer.'),
            ('M', 'Existing system analysis',   'The system shall provide an offline processing capability for reasoning workloads under restricted external connectivity.'),
            ('M', 'Discussion',                 'The system shall treat audit-log writes as best-effort and continue operation when the audit subsystem is unavailable.'),
            ('M', 'Discussion',                 'The system shall treat in-app notifications as best-effort and continue operation when the notification subsystem is unavailable.'),
            ('M', 'Research paper',             'The system shall support idempotent re-ingestion of documents via content hashing, recovering from transient ingestion failures without data loss.'),
        ],
    ),
    (
        'Maintainability',
        [
            ('M', 'Technical feasibility study','The system shall be organised into loosely coupled modules with clear interface boundaries (ingestion, retrieval, reasoning, web).'),
            ('M', 'Technical feasibility study','The system shall pin Python dependencies via a versioned requirements file for reproducible builds.'),
            ('H', 'Existing system analysis',   'The system shall provide an automated test suite covering authentication, role-based access, audit logging, monitoring, and retrieval quality.'),
            ('M', 'Research paper',             'The system shall isolate prompt templates from application code to allow prompt iteration without code changes.'),
            ('M', 'Technical feasibility study','The system shall support swappable LLM providers via configuration, including managed and locally-hosted options.'),
        ],
    ),
    (
        'Usability',
        [
            ('H', 'BBK consultation',           'The system shall display live progress for any background operation expected to exceed five seconds.'),
            ('M', 'Existing system analysis',   'The system shall warn users of impending idle-session timeout and allow them to extend the session without losing the current page state.'),
            ('H', 'Meeting with BBK',           'The system shall provide audit-quality export of analysis reports in XLSX, DOCX, and PDF formats.'),
        ],
    ),
    (
        'Configurability',
        [
            ('M', 'Technical feasibility study','The system shall expose runtime configuration via environment variables for session timeout, MFA enforcement, retry limits, geo-fencing, and LLM provider selection.'),
            ('M', 'Discussion',                 'The system shall expose monitoring thresholds via centralised configuration, including drift windows, bias gaps, and retraining floors.'),
        ],
    ),
    (
        'Observability',
        [
            ('L', 'Research paper',             'The system shall provide optional OpenTelemetry tracing integration that operates as a no-op when not configured.'),
            ('L', 'Research paper',             'The system shall provide optional LangSmith run-logging integration that operates as a no-op when no API key is configured.'),
            ('M', 'Meeting with NAIRDC',        'The system shall provide an administrator monitoring dashboard summarising LLM provider health, vector store statistics, lexical index statistics, active sessions, audit log volume, and recent reasoning failures.'),
        ],
    ),
    (
        'AI Quality & Governance',
        [
            ('M', 'Research paper',             'The system shall evaluate whether AI-generated outputs are semantically supported by the retrieved evidence.'),
            ('M', 'Meeting with NAIRDC',        'The system shall monitor reasoning performance against a rolling baseline, including approval rate, confidence, and validation-error rate.'),
            ('M', 'Meeting with NAIRDC',        'The system shall monitor bias parity across user roles and jurisdictions, flagging approval-rate gaps above a configurable threshold.'),
            ('M', 'Meeting with NAIRDC',        'The system shall surface anomaly signals to administrators, including output distribution skew, validation-error spikes, and failed-login surges.'),
            ('M', 'Meeting with NAIRDC',        'The system shall surface retraining triggers to administrators when reasoning performance crosses configurable thresholds.'),
        ],
    ),
]


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def _set_run_style(run, *, bold=False, color=NAVY, size=10):
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.size = Pt(size)


def _style_paragraph(paragraph, *, bold=False, color=NAVY, size=10, align=None):
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        _set_run_style(run, bold=bold, color=color, size=size)


def main():
    base = Path(__file__).resolve().parent.parent
    out_path = base / 'thesis_non_functional_requirements.docx'
    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = base / f'thesis_non_functional_requirements_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    title = doc.add_heading('Non-Functional Requirements', level=0)
    for run in title.runs:
        run.font.color.rgb = NAVY

    intro = doc.add_paragraph(
        "This section presents the system's non-functional requirements, organized "
        'across eight quality categories: Performance Efficiency, Scalability, '
        'Reliability & Availability, Maintainability, Usability, Configurability, '
        'Observability, and AI Quality & Governance. Each requirement records its '
        'primary elicitation source and a priority level.'
    )
    for run in intro.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY

    total_rows = 1 + sum(1 + len(rows) for _, rows in CATEGORIES)
    table = doc.add_table(rows=total_rows, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    widths = [Cm(1.2), Cm(1.6), Cm(3.4), Cm(10.8)]
    for col_idx, width in enumerate(widths):
        for cell in table.columns[col_idx].cells:
            cell.width = width

    # Header row.
    header = table.rows[0]
    header.cells[0].text = 'ID'
    header.cells[1].text = 'Priority'
    header.cells[2].text = 'Source'
    header.cells[3].text = 'Requirement Description'
    for cell in header.cells:
        shade_cell(cell, '002583')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            _set_run_style(run, bold=True, color=WHITE, size=10)

    nfr_counter = 0
    row_idx = 1
    for category_name, rows in CATEGORIES:
        # Category divider row: merge all four cells, navy fill, white bold text.
        cat_row = table.rows[row_idx]
        merged = (
            cat_row.cells[0]
            .merge(cat_row.cells[1])
            .merge(cat_row.cells[2])
            .merge(cat_row.cells[3])
        )
        merged.text = category_name
        shade_cell(merged, '002583')
        merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = merged.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            _set_run_style(run, bold=True, color=WHITE, size=11)
        row_idx += 1

        for priority, source, desc in rows:
            nfr_counter += 1
            data_row = table.rows[row_idx]
            data_row.cells[0].text = f'NFR-{nfr_counter:02d}'
            data_row.cells[1].text = priority
            data_row.cells[2].text = source
            data_row.cells[3].text = desc
            for cell in data_row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                for paragraph in cell.paragraphs:
                    _style_paragraph(paragraph, color=NAVY, size=10)
            # Bold the ID and the Priority.
            for run in data_row.cells[0].paragraphs[0].runs:
                run.font.bold = True
            for run in data_row.cells[1].paragraphs[0].runs:
                run.font.bold = True
            row_idx += 1

    doc.save(str(out_path))
    print(f'Wrote: {out_path} ({nfr_counter} requirements across {len(CATEGORIES)} categories)')


if __name__ == '__main__':
    main()
