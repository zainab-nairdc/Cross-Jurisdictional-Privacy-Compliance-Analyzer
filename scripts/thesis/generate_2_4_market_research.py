"""Generate §2.4 Market Research as a docx.

Output: thesis_docs/2_4_market_research.docx

Rubric coverage:
  - Lists 5 commercial solutions (OneTrust, TrustArc, BigID, Osano,
    Securiti.ai)
  - Critically discusses each
  - Comparative analysis table covering platform, deployment,
    pricing, strength, missing features
  - Justifies the decision to proceed with CJPCA
  - Verified vendor sources cited

Word count: prose only <= 300 words. Table and references excluded.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_run_style(run, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 1', 2: 'Heading 2'}
    size_map = {1: 16, 2: 14}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  space_after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_caption(doc, label: str, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r1 = p.add_run(label + ' ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run(caption)
    set_run_style(r2, italic=True, color=MUTED, size=10)


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_style(run, color=DARK_GRAY, size=10.5)


def add_table(doc, headers, rows, widths_cm, body_size=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col_idx, w in enumerate(widths_cm):
        for cell in table.columns[col_idx].cells:
            cell.width = Cm(w)
    header_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = h
        shade_cell(cell, '002583')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            set_run_style(run, bold=True, color=WHITE, size=10)
    for r_idx, row_values in enumerate(rows, start=1):
        row = table.rows[r_idx]
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_style(
                        run,
                        bold=(c_idx == 0),
                        color=DARK_GRAY,
                        size=body_size,
                    )


COMPARISON_ROWS = [
    ('OneTrust',
     'SaaS only (no on-premises)',
     'Custom quote, reported US$50,000-300,000+ annually',
     'Curated regulatory library (DataGuidance), broad operational '
     'coverage across 50+ frameworks',
     'No automated cross-jurisdiction clause diff with citations, '
     'regulatory text is reference content for human analysts '
     'rather than input to a diff engine'),
    ('TrustArc',
     'SaaS only (no on-premises)',
     'Custom quote, reported US$10,000-137,000 annually',
     '52,000+ article regulatory research library, assessment '
     'templates, DSR and consent management',
     'No AI-driven clause-level comparison between regulations, '
     'treats regulatory text as curated reference material rather '
     'than as input to an automated diff'),
    ('BigID',
     'SaaS or on-premises',
     'Custom quote, premium-priced',
     'Strongest data discovery and classification, supports '
     'on-premises deployment, multi-language data coverage',
     'Data-side focus only, answers where data lives and who can '
     'reach it rather than how regulatory obligations compare'),
    ('Osano',
     'SaaS only (no on-premises)',
     'US$99-199 per month, free tier available',
     'Transparent SMB pricing, pre-built rule sets for 95+ '
     'regulations, consent banner and DSAR intake',
     'Website and consent-focused, no enterprise legal '
     'interpretation, not for on-premises banking deployment'),
    ('Securiti.ai',
     'SaaS only (no on-premises)',
     'Custom quote',
     'Explicit Bahrain PDPL and India DPDP solution pages, '
     'PrivacyOps engine',
     'Operational compliance focus, marketed as a privacy-ops '
     'platform rather than a citation-backed clause-diff engine'),
]


REFERENCES = [
    'BigID. (n.d.). Data security platform. Retrieved from '
    'https://bigid.com/',

    'OneTrust. (n.d.). Pricing and packaging. Retrieved from '
    'https://www.onetrust.com/pricing/',

    'Osano. (n.d.). Plans and pricing. Retrieved from '
    'https://www.osano.com/plans',

    'Securiti. (n.d.). Bahrain Personal Data Protection Law '
    'solution. Retrieved from https://securiti.ai/solutions/'
    'bahrain-pdpl/',

    'Sprinto. (2025). Honest OneTrust review 2025: Features, '
    'pricing, and alternatives. Retrieved from '
    'https://sprinto.com/blog/onetrust-review/',

    'TrustArc. (n.d.). Privacy management platform. Retrieved from '
    'https://trustarc.com/',
]


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / '2_4_market_research.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'2_4_market_research_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_heading(doc, '2.4 Market Research', level=1)

    add_paragraph(
        doc,
        'The privacy-technology market is well populated with '
        'platforms that automate operational compliance tasks such '
        'as data subject access request handling, consent capture, '
        'data discovery, and vendor risk monitoring. However, no '
        'leading commercial product publicly demonstrates structured, '
        'citation-grounded comparison of legal obligations across '
        'jurisdictions. Four established offerings dominate the '
        'segment, OneTrust (n.d.), TrustArc (n.d.), BigID (n.d.), '
        'and Osano (n.d.), with Securiti (n.d.) as a recent entrant '
        'marketing Bahrain PDPL and India DPDP coverage. Table 12 '
        'at the end of this section summarises the comparative '
        'analysis across platform, deployment model, pricing, key '
        'strengths, and the features missing for cross-jurisdictional '
        'legal reasoning.',
    )

    add_paragraph(
        doc,
        'A consistent pattern emerges across the five platforms. '
        'Their value proposition lies in automating the operational '
        'dimension of privacy compliance, while regulatory '
        'interpretation and cross-jurisdictional comparison remain '
        'manual, consultant-led activities supported by curated '
        'reference libraries rather than automated clause-diff '
        'engines. Independent reviewers note significant '
        'implementation complexity and high total cost of ownership '
        'for enterprise deployments (Sprinto, 2025). Four of the '
        'five are SaaS-only, which can conflict with the '
        'data-residency preferences of Gulf banks operating under '
        'Central Bank of Bahrain oversight. Enterprise pricing also '
        'assumes multi-seat deployment rather than a single-bank '
        'pilot.',
    )

    add_paragraph(
        doc,
        'CJPCA is positioned as an in-house alternative that fills '
        'these gaps without competing on the operational features '
        'the commercial platforms already do well. Built for '
        'on-premises deployment within BBK\'s own infrastructure, '
        'the system removes the cross-border data transfer question '
        'that constrains SaaS platforms in financial-sector '
        'deployment, and replaces per-seat licensing with a '
        'bank-owned platform. The architecture is fully '
        'customisable to BBK\'s specific compliance workflow, '
        'jurisdiction scope, term dictionary, and internal policy '
        'corpus, enabling tailored operation rather than '
        'configuration of a generic product. CJPCA addresses the '
        'interpretive layer that operational platforms leave to '
        'manual effort, complementing rather than competing with '
        'them.',
    )

    add_caption(
        doc, 'Table 12.',
        'Comparative analysis of commercial privacy compliance '
        'platforms.')
    add_table(
        doc,
        headers=['Platform', 'Deployment', 'Pricing', 'Strengths',
                 'Missing for CJPCA use case'],
        rows=COMPARISON_ROWS,
        widths_cm=[2.2, 2.4, 3.4, 4.0, 4.5],
        body_size=9,
    )

    add_heading(doc, 'References for §2.4', level=2)
    for ref in REFERENCES:
        add_reference(doc, ref)

    doc.save(out_path)
    print(f'Wrote {out_path}')


if __name__ == '__main__':
    main()
