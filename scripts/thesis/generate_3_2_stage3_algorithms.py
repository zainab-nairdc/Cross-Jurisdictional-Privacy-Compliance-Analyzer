"""Generate Section 3.2.3 Algorithms + Appendix 2.7 as a Word document.

Output: thesis_docs/3_2_3_algorithms_full.docx

Contents (main body):
  3.2.3 Algorithms
    3.2.3.1 Algorithm notation in use
    3.2.3.2 System overview                 + Figure 8
    3.2.3.3 The reasoning agent             + Alg 1 + Alg 1a + Figure 9
    3.2.3.4 The agent's input pipeline      + Alg 2 + Figure 10
    3.2.3.5 Workflow: Policy mapping        + Alg 3 + Figure 11
    3.2.3.6 Workflow: Comparison            + Alg 4 + Figure 12
    3.2.3.7 Cross-cutting: ingestion,
            audit, export                   + Alg 5 + Alg 6 + Figure 13 + Figure 14

Contents (appendix):
  Appendix 2.7 Algorithm reference
    A.7.1 Auxiliary algorithms (9 items, a–i)
    A.7.2 Django integration (3 items, a–c)
    A.7.3 Lookup tables — reference to separate docx

Style follows the Stage 1 + Stage 2 generators (navy headings, 2 cm margins).
Embeds the PNGs at their actual filenames from the user's screenshot workflow.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
DARK_GRAY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def set_run_style(run, *, bold=False, italic=False, color=NAVY, size=11):
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_paragraph(doc, text, *, bold=False, italic=False, color=DARK_GRAY,
                  size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_style(run, bold=bold, italic=italic, color=color, size=size)
    return p


def add_heading(doc, text, *, level=1):
    style_map = {1: 'Heading 2', 2: 'Heading 3', 3: 'Heading 4'}
    size_map = {1: 14, 2: 12, 3: 11}
    h = doc.add_paragraph(style=style_map[level])
    run = h.add_run(text)
    run.font.color.rgb = NAVY
    run.font.bold = True
    run.font.size = Pt(size_map[level])
    return h


def add_figure_caption(doc, label: str, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(14)
    r1 = p.add_run(label + ' ')
    set_run_style(r1, bold=True, color=NAVY, size=10)
    r2 = p.add_run(caption)
    set_run_style(r2, italic=True, color=MUTED, size=10)


def add_image(doc, image_path: Path, width_cm: float = 16.0):
    if not image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'[MISSING IMAGE: {image_path.name}]')
        set_run_style(run, bold=True, italic=True, color=MUTED, size=10)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def main():
    base = Path(__file__).resolve().parents[2]
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / '3_2_3_algorithms_full.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = out_dir / f'3_2_3_algorithms_full_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break

    diagrams = base / 'diagrams'
    pseudo   = base / 'thesis_docs' / 'pseudocode'

    # ── Flowchart PNG paths (the user's actual filenames) ─────────────────
    fig_8  = diagrams / 'CJPCA Master System Flowchart.png'
    fig_9  = diagrams / 'Figure_9_ReasoningAgent.png'
    fig_10 = diagrams / 'Figure_10_HybridRetrieval.png'
    fig_11 = diagrams / 'Policy Mapping Workflow.png'
    fig_12 = diagrams / 'Cross-Jurisdictional Comparison Workflow.png'
    fig_13 = diagrams / 'Figure_13_Ingestion.png'
    fig_14 = diagrams / 'Approved-Report Export with Audit Hash.png'

    # ── Algorithm pseudocode PNG paths ────────────────────────────────────
    alg_1   = pseudo / 'reasoing agent.png'
    alg_1a  = pseudo / 'vertify and score.png'
    alg_2   = pseudo / 'Algorithm 2 — HybridRetrieve.png'
    alg_3   = pseudo / 'Algorithm 3 — RunPolicyMapping.png'
    alg_4   = pseudo / 'Algorithm 4 — RunComparison.png'
    alg_5   = pseudo / 'Algorithm 5 — IngestDocument.png'
    alg_6   = pseudo / 'Algorithm 6 — AuditHashAndExport.png'

    a_7_1_a = pseudo / 'Algorithm A.7.1.a — StrictnessScore.png'
    a_7_1_b = pseudo / 'Algorithm A.7.1.b — DivergenceRanking.png'
    a_7_1_c = pseudo / 'riskscore.png'
    a_7_1_d = pseudo / 'query router.png'
    a_7_1_e = pseudo / 'termdicexpan.png'
    a_7_1_f = pseudo / 'autoroutescope.png'
    a_7_1_g = pseudo / 'chunkclassift.png'
    a_7_1_h = pseudo / 'safe fallback.png'
    a_7_1_i = pseudo / 'local llm.png'

    a_7_2_a = pseudo / 'sync_gap_on_approval.png'   # may be missing
    a_7_2_b = pseudo / 'launch_subprocess.png'      # may be missing
    a_7_2_c = pseudo / 'cascade approval.png'

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # ── 3.2.3 Algorithms ──────────────────────────────────────────────────
    add_heading(doc, '3.2.3 Algorithms', level=1)

    # 3.2.3.1 Algorithm notation in use  ~50 words
    add_heading(doc, '3.2.3.1 Algorithm notation in use', level=2)
    add_paragraph(
        doc,
        'Each algorithm is presented as numbered pseudocode with an Input '
        'and Output declaration. Assignments use the arrow ←, comparisons '
        'use plain operators, and comments use //. Cross-references to '
        'other algorithms appear as inline pointers. Auxiliary procedures '
        'used by the main-body algorithms are expanded in Appendix 2.7.',
    )

    # 3.2.3.2 System overview  ~75 words
    add_heading(doc, '3.2.3.2 System overview', level=2)
    add_paragraph(
        doc,
        'Figure 8 puts the entire system on one page as a flowchart. '
        'After authentication and a role check, an authenticated session '
        'branches into one of five user actions: an administrator uploads '
        'a document, an analyst runs a policy mapping or a comparison, an '
        'analyst or reviewer asks the Copilot, or a reviewer validates '
        'pending items. Each branch eventually writes one or more rows to '
        'the audit log, so every flow ends with traceability.',
    )
    add_image(doc, fig_8, width_cm=15.0)
    add_figure_caption(doc, 'Figure 8.', 'CJPCA Master System Flowchart')

    # 3.2.3.3 The reasoning agent
    add_heading(doc, '3.2.3.3 The reasoning agent', level=2)
    add_paragraph(
        doc,
        'The reasoning layer is a LangGraph state machine, referred to as '
        'the reasoning agent. Algorithm 1 shows its main loop. A query is '
        'routed by the 3-NN classifier in A.7.1.d; the agent then drafts '
        'a structured answer and verifies it (Algorithm 1a). Verification '
        'combines citation grounding with cross-chunk recovery and an NLI '
        'hallucination score that takes the MAX entailment across the '
        'top-5 chunks. MAX (not mean) is used so a multi-sentence answer '
        'is judged grounded if any chunk supports it. A failed '
        'verification re-prompts with the verifier’s complaint; once '
        'retries exhaust, the agent returns a deterministic safe answer. '
        'All LLM calls go to the local Ollama daemon on localhost:11434 '
        '(default model llama3.2:1b); the HTTP call and OutputFixingParser '
        'retry are in Algorithm A.7.1.i.',
    )
    add_image(doc, alg_1, width_cm=15.0)
    add_figure_caption(doc, 'Algorithm 1.', 'ReasoningAgent')
    add_image(doc, alg_1a, width_cm=15.0)
    add_figure_caption(doc, 'Algorithm 1a.', 'VerifyAndScore')
    add_image(doc, fig_9, width_cm=14.0)
    add_figure_caption(doc, 'Figure 9.', 'Reasoning Agent State Machine')

    # 3.2.3.4 The agent's input pipeline
    add_heading(doc, '3.2.3.4 The agent’s input pipeline', level=2)
    add_paragraph(
        doc,
        'The agent does not retrieve evidence itself; that work is done '
        'by the hybrid retrieval pipeline in Algorithm 2 and Figure 10. '
        'A query is first expanded against the cross-jurisdictional term '
        'dictionary (A.7.1.e), then drives two parallel searches: BM25 '
        'over SQLite FTS5, and vector search over ChromaDB. BM25 catches '
        'the legal acronyms (DPO, PDPL, GDPR) that vector embeddings '
        'miss; the vector search catches paraphrases that BM25 misses. '
        'The two lists are merged with Reciprocal Rank Fusion (k = 60), '
        'then the top-20 candidates are reranked by a cross-encoder '
        'before the top-5 are returned.',
    )
    add_image(doc, alg_2, width_cm=15.0)
    add_figure_caption(doc, 'Algorithm 2.', 'HybridRetrieve')
    add_image(doc, fig_10, width_cm=14.0)
    add_figure_caption(doc, 'Figure 10.', 'Hybrid Retrieval Pipeline')

    # 3.2.3.5 Workflow: Policy mapping  ~95 words
    add_heading(doc, '3.2.3.5 Workflow: Policy mapping', level=2)
    add_paragraph(
        doc,
        'Algorithm 3 traces a policy mapping from analyst click to '
        'populated rows. The view creates a parent MappingAnalysis row, '
        'dispatches a background subprocess (A.7.2.b), and returns '
        'immediately. Inside the subprocess, auto-routing (A.7.1.f) '
        'reduces the obligation set to topics the policy actually covers '
        'when scope_mode = AUTO. The subprocess then iterates each '
        'remaining obligation: hybrid retrieval finds evidence, the '
        'reasoning agent verifies an answer, the confidence rule caps '
        'unverified rows at 0.6, and a post-save signal (A.7.2.a) '
        'reconciles the Gap table. Figure 11 traces the full flow.',
    )
    add_image(doc, alg_3, width_cm=15.0)
    add_figure_caption(doc, 'Algorithm 3.', 'RunPolicyMapping')
    add_image(doc, fig_11, width_cm=14.0)
    add_figure_caption(doc, 'Figure 11.', 'Policy Mapping Workflow')

    # 3.2.3.6 Workflow: Cross-jurisdictional comparison
    add_heading(doc, '3.2.3.6 Workflow: Cross-jurisdictional comparison',
                level=2)
    add_paragraph(
        doc,
        'Algorithm 4 traces a comparison from click to populated rows '
        'using the same dispatch pattern. Inside the subprocess the '
        'strictness score (A.7.1.a) is computed once per regulation as '
        'a four-factor weighted formula. Then for each topic, hybrid '
        'retrieval pulls chunks from both regulations in parallel and '
        'the reasoning agent classifies the relationship as equivalent, '
        'stricter, additional, or conflicting. After all topics finish, '
        'the divergence-ranking formula (A.7.1.b) orders rows for the '
        'executive summary, with conflicting clauses surfaced before '
        'lesser divergences. Figure 12 traces the flow.',
    )
    add_image(doc, alg_4, width_cm=15.0)
    add_figure_caption(doc, 'Algorithm 4.', 'RunComparison')
    add_image(doc, fig_12, width_cm=14.0)
    add_figure_caption(doc, 'Figure 12.',
                       'Cross-Jurisdictional Comparison Workflow')

    # 3.2.3.7 Cross-cutting: ingestion, audit, export
    add_heading(doc, '3.2.3.7 Cross-cutting: ingestion, audit, and export',
                level=2)
    add_paragraph(
        doc,
        'Two algorithms surround the agent. Algorithm 5 (ingestion) '
        'extracts text, OCRs scanned PDFs, applies two-level chunking '
        '(header split then semantic merge), embeds and classifies '
        'each chunk, scans for prompt-injection content via Tier A '
        'regex plus Tier B LLM judge, then indexes safe chunks into '
        'ChromaDB and SQLite FTS5; flagged chunks go to quarantine for '
        'admin review. Algorithm 6 (export) computes the audit-hash '
        'fingerprint, sorts rows so reviewer-touched entries surface '
        'first, applies a severity-driven due-date cascade, and renders '
        'the PDF or XLSX. The 16-character SHA-256 prefix makes silent '
        'post-signing edits detectable. Figures 13 and 14 trace both '
        'flows.',
    )
    add_image(doc, alg_5, width_cm=15.0)
    add_figure_caption(doc, 'Algorithm 5.', 'IngestDocument')
    add_image(doc, alg_6, width_cm=15.0)
    add_figure_caption(doc, 'Algorithm 6.', 'AuditHashAndExport')
    add_image(doc, fig_13, width_cm=15.0)
    add_figure_caption(doc, 'Figure 13.', 'Ingestion Pipeline')
    add_image(doc, fig_14, width_cm=14.0)
    add_figure_caption(doc, 'Figure 14.',
                       'Approved-Report Export with Audit Hash')

    # ── PAGE BREAK → Appendix 2.7 ────────────────────────────────────────
    add_page_break(doc)

    # ── Appendix 2.7 ──────────────────────────────────────────────────────
    add_heading(doc, 'Appendix 2.7 — Algorithm reference', level=1)
    add_paragraph(
        doc,
        'This appendix expands the algorithms named but not boxed in '
        '§3.2.3. A.7.1 covers auxiliary procedures inside the AI '
        'pipeline (scoring, classification, retrieval helpers, the LLM '
        'invocation). A.7.2 covers the three Django integration '
        'algorithms that connect the AI pipeline to the web layer. '
        'A.7.3 lists the configuration constants, weights, and lookup '
        'tables the algorithms read from; that section is provided as '
        'a separate document (appendix_2_7_lookups.docx) due to size.',
    )

    # A.7.1 Auxiliary algorithms
    add_heading(doc, 'A.7.1 Auxiliary algorithms', level=2)
    add_paragraph(
        doc,
        'Nine procedures expanded here. Each carries the formula or '
        'logic that the main-body algorithms reference by name. '
        'Weights, thresholds, and lookup tables are in A.7.3.',
    )

    for (label, caption, image) in [
        ('Algorithm A.7.1.a.', 'StrictnessScore',          a_7_1_a),
        ('Algorithm A.7.1.b.', 'DivergenceRanking',        a_7_1_b),
        ('Algorithm A.7.1.c.', 'RiskScore',                a_7_1_c),
        ('Algorithm A.7.1.d.', 'QueryRouter',              a_7_1_d),
        ('Algorithm A.7.1.e.', 'TermDictionaryExpand',     a_7_1_e),
        ('Algorithm A.7.1.f.', 'AutoRouteScope',           a_7_1_f),
        ('Algorithm A.7.1.g.', 'ChunkClassify',            a_7_1_g),
        ('Algorithm A.7.1.h.', 'SafeFallback',             a_7_1_h),
        ('Algorithm A.7.1.i.', 'InvokeLocalLLM',           a_7_1_i),
    ]:
        add_image(doc, image, width_cm=15.0)
        add_figure_caption(doc, label, caption)

    # A.7.2 Django integration
    add_heading(doc, 'A.7.2 Django integration algorithms', level=2)
    add_paragraph(
        doc,
        'Three patterns that glue the AI pipeline to the Django web '
        'layer. Each illustrates a distinct integration mechanism: '
        'a Django signal cascade, an async dispatch pattern, and a '
        'manual cross-app cascade.',
    )

    for (label, caption, image) in [
        ('Algorithm A.7.2.a.', 'SyncGapOnApproval', a_7_2_a),
        ('Algorithm A.7.2.b.', 'LaunchSubprocess',  a_7_2_b),
        ('Algorithm A.7.2.c.', 'CascadeApproval',   a_7_2_c),
    ]:
        add_image(doc, image, width_cm=15.0)
        add_figure_caption(doc, label, caption)

    # A.7.3 reference
    add_heading(doc, 'A.7.3 Lookup tables', level=2)
    add_paragraph(
        doc,
        'Eight reference tables (chunk_tags taxonomy, per-jurisdiction '
        'risk weights, per-topic impact weights, coverage factors, '
        'severity-to-due-date buckets, algorithm configuration '
        'constants, DivergenceRanking weights, retrieval evaluation '
        'benchmarks) are in the separate document '
        'appendix_2_7_lookups.docx and should be appended after this '
        'section in the final thesis.',
    )

    doc.save(out_path)
    print(f'Wrote: {out_path}')


if __name__ == '__main__':
    main()
