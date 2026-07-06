"""Combine §3.1 Requirements main body and Appendix 1 into one docx.

Output: thesis_docs/3_1_requirements_combined.docx

Imports data and writer functions from:
  - generate_3_1_requirements.py (main body)
  - generate_appendix_1_requirements.py (appendix)
"""

import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def main():
    base = Path(__file__).resolve().parents[2]
    scripts_dir = base / 'scripts' / 'thesis'
    out_dir = base / 'thesis_docs'
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / '3_1_requirements_combined.docx'

    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = (out_dir /
                             f'3_1_requirements_combined_v{n}.docx')
                if not candidate.exists():
                    out_path = candidate
                    break

    mb = load_module('mb_requirements',
                     scripts_dir / 'generate_3_1_requirements.py')
    ap = load_module('ap_requirements',
                     scripts_dir / 'generate_appendix_1_requirements.py')

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # ---- Main body §3.1 ----
    mb.add_heading(doc, '3.1 Requirements', level=1)
    mb.add_paragraph(
        doc,
        'This section describes the requirements gathered for CJPCA. '
        'It covers the elicitation methodology, user-voice '
        'requirements grouped by operational role, twenty functional '
        'requirements split across six categories, twenty '
        'non-functional requirements split across seven quality '
        'categories, and the six development constraints that '
        'shaped scope. Supporting evidence (meeting records, '
        'research findings, document analysis, and the technical '
        'feasibility study) appears in Appendix 1.',
    )
    mb.write_311(doc)
    mb.write_312(doc)
    mb.write_313(doc)
    mb.write_314(doc)
    mb.write_315(doc)

    # ---- Page break before appendix ----
    doc.add_page_break()

    # ---- Appendix 1 ----
    ap.add_heading(doc, 'Appendix 1 — Requirements', level=1)
    ap.add_paragraph(
        doc,
        'This appendix records the underlying evidence gathered '
        'during the requirements-elicitation phase referenced in '
        '§3.1, organised by elicitation technique (stakeholder '
        'meetings, research, document analysis, technical '
        'feasibility study), and concludes with the extended '
        'functional requirements catalogue.',
    )
    ap.write_a11(doc)
    ap.write_a12(doc)
    ap.write_a13(doc)
    ap.write_a14(doc)
    ap.write_a15(doc)
    ap.write_a16(doc)
    ap.write_references(doc)

    doc.save(out_path)
    print(f'Wrote {out_path}')


if __name__ == '__main__':
    main()
