"""Generate the Functional Requirements Word document for the thesis.

Output: thesis_functional_requirements.docx in the project root.

Layout: one continuous table with four columns
(ID, Priority, Source, Requirement Description). Phase names appear as
merged-cell section headers between requirement groups.
Priority: H (must-have), M (should-have), L (could-have) — MoSCoW.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x00, 0x25, 0x83)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


PHASES = [
    (
        'Document Ingestion',
        [
            ('H', 'BBK consultation',          'The system shall ingest privacy regulations from Bahrain, India, and Kuwait.'),
            ('H', 'Brief',                     "The system shall ingest BBK's internal privacy policies, governance documents, and operational SOPs."),
            ('H', 'Legal document analysis',   'The system shall store metadata for every document, including jurisdiction, version, publication date, and source.'),
            ('H', 'Technical feasibility study', 'The system shall convert PDF, DOCX, and HTML documents into clean text for processing.'),
            ('H', 'Research paper',            'The system shall split documents into searchable chunks that preserve their section structure.'),
            ('M', 'Technical feasibility study', 'The system shall tag each chunk by its legal-text type (operative rule, definition, preamble, or general).'),
            ('H', 'Research paper',            'The system shall make every chunk searchable by meaning (semantic) and by keyword (full-text).'),
            ('M', 'Discussion',                'The system shall process ingestion in the background with live progress reporting.'),
            ('H', 'Brief',                     'The system shall provide an administrative interface for document uploads.'),
        ],
    ),
    (
        'Retrieval Layer',
        [
            ('H', 'Research paper',            'The system shall retrieve and rank chunks using combined semantic and keyword search.'),
            ('M', 'Research paper',            'The system shall rerank retrieved results to improve relevance before analysis.'),
            ('H', 'Research paper',            'The system shall expand queries using equivalent legal terminology across jurisdictions.'),
            ('M', 'Discussion',                'The system shall filter retrieval results by jurisdiction, topic, subcategory, and document title.'),
            ('M', 'Technical feasibility study', 'The system shall preserve parent-child chunk relationships during retrieval.'),
            ('H', 'Legal document analysis',   'The system shall return citation-preserving results linked to their source documents.'),
        ],
    ),
    (
        'Reasoning & Analysis',
        [
            ('H', 'Meeting with BBK',         'The system shall compare obligations between two selected regulatory frameworks for a chosen compliance topic.'),
            ('H', 'Legal document analysis',  'The system shall identify similarities, differences, and conflicts between compared regulatory obligations.'),
            ('H', 'Legal document analysis',  'The system shall determine equivalence levels between obligations across jurisdictions.'),
            ('M', 'Legal document analysis',  'The system shall assess comparative strictness between regulatory frameworks across procedural, substantive, and enforcement dimensions.'),
            ('H', 'Meeting with BBK',         'The system shall map internal policy documents against one or more regulations.'),
            ('H', 'BBK consultation',         'The system shall classify policy coverage as Fully Covered, Partially Covered, Requires Review, or Not Covered.'),
            ('M', 'Meeting with NAIRDC',      'The system shall generate remediation recommendations for uncovered obligations and identified compliance gaps.'),
            ('M', 'Legal document analysis',  'The system shall detect compliance topics contained within uploaded policy documents.'),
            ('M', 'Legal document analysis',  'The system shall classify document chunks into predefined compliance topics (such as consent, breach notification, retention, and cross-border transfer) at ingest time and on demand during workflow execution.'),
            ('H', 'Discussion',               'The system shall automatically route policy mapping workflows to relevant regulations based on detected compliance topics.'),
            ('H', 'Meeting with BBK',         'The system shall aggregate uncovered obligations from multiple jurisdictions into a unified gap analysis report.'),
            ('H', 'BBK consultation',         'The system shall prioritise compliance gaps using severity levels of High, Medium, and Low.'),
            ('M', 'Discussion',               'The system shall return partially verified results when full citation verification cannot be completed.'),
            ('M', 'Legal document analysis',  'The system shall maintain a dictionary of equivalent legal terminology across supported jurisdictions.'),
            ('H', 'Meeting with BBK',         'The system shall support a strict evidence mode that restricts reasoning to analyst-selected documents only.'),
            ('M', 'Discussion',               'The system shall support an open evidence mode that augments selected documents with related chunks from the broader corpus.'),
        ],
    ),
    (
        'Web Application',
        [
            # Home dashboard
            ('H', 'Meeting with BBK',       'The system shall provide a home dashboard summarising overall compliance posture.'),
            ('H', 'BBK consultation',       'The system shall display overall compliance coverage percentages.'),
            ('H', 'BBK consultation',       'The system shall display compliance coverage breakdowns by jurisdiction.'),
            ('M', 'Meeting with NAIRDC',    'The system shall display pending compliance actions and workflow prompts.'),
            ('M', 'Discussion',             'The system shall display recent analyses and document ingestion activities.'),
            # Library
            ('H', 'Meeting with BBK',       'The system shall provide a library workspace for browsing regulations, internal policies, obligations, and legal terminology dictionaries.'),
            ('H', 'Meeting with BBK',       'The system shall support cross-reference search across the document corpus.'),
            ('M', 'Discussion',             'The system shall support user-defined document tagging.'),
            # Document lifecycle
            ('H', 'Meeting with BBK',       'The system shall support analyst-driven document uploads.'),
            ('M', 'Technical feasibility study', 'The system shall display document metadata previews before ingestion.'),
            ('M', 'BBK consultation',       'The system shall display live ingestion progress during document processing.'),
            ('M', 'Meeting with NAIRDC',    'The system shall support document deletion.'),
            ('H', 'BBK consultation',       'The system shall provide full-text document viewing with citation labels.'),
            # Comparison workspace
            ('H', 'Meeting with BBK',       'The system shall provide a comparison workspace for configuring cross-regulation comparisons.'),
            ('M', 'BBK consultation',       'The system shall display live workflow progress during comparison analysis.'),
            ('H', 'Meeting with BBK',       'The system shall provide comparison results through overview, clauses, equivalency graph, relationship map, topic map, and insights views.'),
            ('M', 'Legal document analysis', 'The system shall display strictness indicators for compared regulatory frameworks.'),
            ('M', 'Meeting with NAIRDC',    'The system shall allow reviewers to transition per-clause statuses and record reviewer notes.'),
            # Policy mapping workspace
            ('H', 'Meeting with BBK',       'The system shall provide a policy mapping workspace for configuring policy-to-regulation mapping workflows.'),
            ('M', 'BBK consultation',       'The system shall display live workflow progress during policy mapping analysis.'),
            ('L', 'Discussion',             'The system shall allow analysts to cancel running policy mapping workflows.'),
            ('H', 'Meeting with BBK',       'The system shall provide per-obligation mapping results with supporting evidence panels.'),
            ('M', 'Meeting with NAIRDC',    'The system shall allow analysts to manually override obligation coverage classifications.'),
            ('M', 'Meeting with NAIRDC',    'The system shall allow reviewers to transition per-obligation statuses and record reviewer notes.'),
            # Remediation
            ('M', 'Meeting with BBK',       'The system shall allow analysts to record manual remediation recommendations for identified compliance gaps.'),
            # Dual mapping
            ('M', 'BBK consultation',       'The system shall provide a dual-mapping view for comparing one policy against two regulations simultaneously.'),
            # Workflow handoff
            ('H', 'BBK consultation',       'The system shall support workflow handoff from analysts to reviewers.'),
            ('H', 'BBK consultation',       'The system shall provide reviewer approval workflows for completed analyses.'),
            # Review queue + export
            ('H', 'BBK consultation',       'The system shall provide a reviewer validation queue for reviewing AI-generated findings.'),
            ('H', 'BBK consultation',       'The system shall allow reviewers to approve, reject, or modify AI-generated findings.'),
            ('H', 'Meeting with BBK',       'The system shall support export of analysis reports in XLSX, DOCX, and PDF formats.'),
            # Gap workspaces
            ('H', 'Meeting with BBK',       'The system shall provide a cross-jurisdiction gap analysis workspace.'),
            ('H', 'BBK consultation',       'The system shall provide a unified compliance gap register.'),
            ('M', 'Legal document analysis', 'The system shall provide a regulatory conflict analysis workspace.'),
            ('M', 'Discussion',             'The system shall allow reviewers and administrators to assign compliance gaps to users with due dates.'),
            # Analytics
            ('M', 'Existing system analysis', 'The system shall provide a role-based analytics dashboard summarising compliance coverage, gaps, conflicts, and reviewer activity.'),
            ('M', 'Existing system analysis', 'The system shall provide interactive compliance coverage heatmaps.'),
            # Audit history
            ('H', 'BBK consultation',       'The system shall maintain an audit history of all mutating system actions.'),
            ('M', 'BBK consultation',       'The system shall provide detailed audit event views.'),
            # Doc viewer + term overlay (shared UI features used by Copilot and the workspaces)
            ('H', 'Legal document analysis', 'The system shall support citation click-through navigation to source documents with highlighted evidence excerpts.'),
            ('M', 'Legal document analysis', 'The system shall provide inline legal terminology translation across jurisdictions.'),
        ],
    ),
    (
        'Copilot Agent',
        [
            ('H', 'Meeting with BBK',         'The system shall provide a global AI Copilot chat overlay accessible from every page.'),
            ('M', 'Discussion',               'The system shall classify each Copilot query by intent (comparison analysis, policy gap, regulatory lookup) and select the appropriate retrieval source.'),
            ('H', 'Research paper',           'The system shall ground every Copilot answer in retrieved corpus chunks before responding.'),
            ('M', 'Legal document analysis',  'The system shall preferentially ground Copilot answers in reviewer-approved analyses, falling back to indexed regulatory text otherwise.'),
            ('L', 'BBK consultation',         'The system shall allow users to opt in to draft (unreviewed) analyses as a grounding source via a per-session toggle.'),
            ('M', 'Meeting with BBK',         'The system shall allow users to constrain Copilot retrieval to a single selected document.'),
            ('H', 'Research paper',           'The system shall route Copilot queries through the same verification pipeline used by structured reasoning workflows.'),
            ('H', 'Research paper',           'The system shall return Copilot answers with regulation-name, article-reference, and excerpt-level citations.'),
            ('M', 'Research paper',           'The system shall report a confidence score and a hallucination risk score alongside every Copilot answer.'),
            ('M', 'Meeting with NAIRDC',      'The system shall compute and surface a Copilot readiness state of not-ready, partially-ready, or ready, based on the availability of regulations, internal policies, approved comparisons, and approved mappings.'),
            ('L', 'Discussion',               'The system shall display a fallback explanation message when Copilot answers from raw regulatory text instead of approved analyses.'),
            ('L', 'Technical feasibility study', 'The system shall maintain per-user Copilot conversation history across page navigation, retaining the most recent twenty messages.'),
            ('L', 'Discussion',               'The system shall allow users to clear Copilot conversation history.'),
            ('M', 'Research paper',           'The system shall degrade gracefully when no relevant chunks are retrieved or when the AI provider is unavailable, returning a clear non-fabricated response.'),
        ],
    ),
    (
        'Security & Access Control',
        [
            # Authentication & MFA
            ('H', 'Existing system analysis',  'The system shall require username and password authentication for every user account.'),
            ('H', 'Existing system analysis',  'The system shall require time-based one-time password (TOTP) multi-factor authentication for every user account.'),
            ('H', 'Existing system analysis',  'The system shall enforce TOTP enrolment before any user can access business functionality after first login.'),
            ('M', 'Existing system analysis',  'The system shall apply multi-factor authentication to the Django administrative interface.'),
            # Password policy
            ('H', 'Existing system analysis',  'The system shall enforce password complexity requirements of minimum twelve characters, including uppercase, lowercase, digit, and symbol classes.'),
            ('M', 'Existing system analysis',  "The system shall reject passwords that contain the user's username or email substring."),
            ('H', 'Existing system analysis',  'The system shall require users to change temporary passwords on first login or after an administrator-initiated reset.'),
            # Brute-force defense
            ('H', 'Existing system analysis',  'The system shall lock user accounts after a configurable threshold of failed login attempts, keyed on both username and source IP address.'),
            # Session security
            ('H', 'Existing system analysis',  'The system shall log users out automatically after a configurable period of inactivity.'),
            ('M', 'Existing system analysis',  'The system shall enforce an absolute session lifetime independent of activity.'),
            ('H', 'Research paper',            'The system shall harden session and CSRF cookies with HttpOnly, SameSite, and Secure flags in production.'),
            ('H', 'Existing system analysis',  'The system shall terminate all active sessions of a user when their role, password, or MFA configuration is changed by an administrator.'),
            # Geo-fencing
            ('L', 'Meeting with BBK',          'The system shall optionally restrict access to a configurable country allowlist using GeoIP lookups, allowing access when the lookup database is unavailable.'),
            # Role-based access control
            ('H', 'BBK consultation',          'The system shall enforce role-based access control with three roles: Compliance Analyst, Legal Reviewer, and Administrator.'),
            ('H', 'Existing system analysis',  'The system shall enforce role checks at every authenticated route and on every privileged action.'),
            # User lifecycle management
            ('H', 'Existing system analysis',  'The system shall not provide self-registration; user accounts shall be created by administrators only.'),
            ('H', 'Existing system analysis',  'The system shall allow administrators to create user accounts and assign one of the three roles (Analyst, Reviewer, Administrator).'),
            ('H', 'Existing system analysis',  'The system shall allow administrators to change user roles, disable accounts, reset passwords, and reset MFA devices.'),
            ('M', 'Existing system analysis',  'The system shall deliver temporary credentials to new users by email and require them to change the password on first login.'),
            ('M', 'Technical feasibility study','The system shall support SMTP-based email delivery for credential and notification messages, configurable via environment variables.'),
            # Web application hardening
            ('H', 'Research paper',            'The system shall enforce Cross-Site Request Forgery (CSRF) protection on every state-changing request.'),
            ('H', 'Research paper',            'The system shall set hardened HTTP security headers including HSTS, X-Frame-Options, X-Content-Type-Options, and Referrer-Policy.'),
            ('H', 'Research paper',            'The system shall automatically redirect HTTP requests to HTTPS in production and enforce a one-year HSTS policy with subdomain coverage.'),
            ('M', 'Research paper',            'The system shall enforce a Content Security Policy with an explicit allowlist of script, style, font, image, and form-action sources.'),
            # Audit logging
            ('H', 'BBK consultation',          'The system shall maintain a global audit log of every security-relevant action, capturing actor, role at time of action, IP address, target object, and metadata.'),
            ('H', 'Existing system analysis',  'The system shall automatically log authentication events (login, logout, login failure, MFA enrolment, idle timeout) without requiring per-view instrumentation.'),
            ('M', 'BBK consultation',          'The system shall maintain a per-result lifecycle audit trail recording every status transition on every comparison and mapping result with actor, from-state, to-state, and JSON diff.'),
            ('M', 'BBK consultation',          'The system shall provide an administrator-only audit log viewer with filtering by category, user, and date range.'),
            # AI/ML prompt-injection defense
            ('H', 'Research paper',            'The system shall scan every ingested chunk for prompt-injection content using a catalogue of known attack patterns.'),
            ('M', 'Research paper',            'The system shall escalate borderline chunks to a second-tier classifier with hardened prompting.'),
            ('H', 'Meeting with NAIRDC',       'The system shall quarantine flagged chunks until an administrator approves or rejects them, recording every decision in the audit log.'),
            # RAG integrity (functional integrity gates)
            ('H', 'Research paper',            'The system shall verify every AI-generated citation against the underlying retrieved chunks before returning results, rejecting any citation whose quoted text does not appear in source.'),
            ('H', 'Research paper',            'The system shall constrain AI outputs to a strict structured schema, automatically repairing and rejecting malformed responses.'),
            # Data isolation (functional)
            ('H', 'Technical feasibility study', 'The system shall enforce document-scope filtering at the database layer so that the language model physically cannot access chunks outside the analyst-selected scope.'),
        ],
    ),
]


def shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def _set_run_style(run, *, bold=False, color=NAVY, size=10):
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.size = Pt(size)


def _style_paragraph(paragraph, *, bold=False, color=NAVY, size=10, align=None):
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        _set_run_style(run, bold=bold, color=color, size=size)


def main():
    base = Path(__file__).resolve().parent.parent
    out_path = base / 'thesis_functional_requirements.docx'
    # If the file is locked (open in Word), fall back to a versioned filename.
    if out_path.exists():
        try:
            with open(out_path, 'ab'):
                pass
        except PermissionError:
            for n in range(2, 100):
                candidate = base / f'thesis_functional_requirements_v{n}.docx'
                if not candidate.exists():
                    out_path = candidate
                    break
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    title = doc.add_heading('Functional Requirements', level=0)
    for run in title.runs:
        run.font.color.rgb = NAVY

    intro = doc.add_paragraph(
        "This section presents the system's functional requirements, organized "
        'across six pipeline phases: Document Ingestion, Retrieval Layer, '
        'Reasoning & Analysis, Web Application, Copilot Agent, and Security & '
        'Access Control. Each requirement records its primary elicitation source '
        'and a priority level.'
    )
    for run in intro.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY

    total_rows = 1 + sum(1 + len(rows) for _, rows in PHASES)
    table = doc.add_table(rows=total_rows, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    widths = [Cm(1.2), Cm(1.6), Cm(3.4), Cm(10.8)]
    for col_idx, width in enumerate(widths):
        for cell in table.columns[col_idx].cells:
            cell.width = width

    # Header row.
    header = table.rows[0]
    header.cells[0].text = 'ID'
    header.cells[1].text = 'Priority'
    header.cells[2].text = 'Source'
    header.cells[3].text = 'Requirement Description'
    for cell in header.cells:
        shade_cell(cell, '002583')
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            _set_run_style(run, bold=True, color=WHITE, size=10)

    fr_counter = 0
    row_idx = 1
    for phase_name, rows in PHASES:
        # Phase divider row: merge all four cells, navy fill, white bold text.
        phase_row = table.rows[row_idx]
        merged = (
            phase_row.cells[0]
            .merge(phase_row.cells[1])
            .merge(phase_row.cells[2])
            .merge(phase_row.cells[3])
        )
        merged.text = phase_name
        shade_cell(merged, '002583')
        merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = merged.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            _set_run_style(run, bold=True, color=WHITE, size=11)
        row_idx += 1

        for priority, source, desc in rows:
            fr_counter += 1
            data_row = table.rows[row_idx]
            data_row.cells[0].text = str(fr_counter)
            data_row.cells[1].text = priority
            data_row.cells[2].text = source
            data_row.cells[3].text = desc
            for cell in data_row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                for paragraph in cell.paragraphs:
                    _style_paragraph(paragraph, color=NAVY, size=10)
            # Bold the ID and the Priority.
            for run in data_row.cells[0].paragraphs[0].runs:
                run.font.bold = True
            for run in data_row.cells[1].paragraphs[0].runs:
                run.font.bold = True
            row_idx += 1

    doc.save(str(out_path))
    print(f'Wrote: {out_path} ({fr_counter} requirements across {len(PHASES)} phases)')


if __name__ == '__main__':
    main()
